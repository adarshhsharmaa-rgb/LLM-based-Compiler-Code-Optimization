"""Builds the Review 3 (final) deck.

It is the Review 2 deck (``build_review2.py`` builds every one of its slides) with the
architecture brought up to date, followed by the Review 3 work: LLM self-repair, execution
time in the cost model, SMT verification, the dataset of real programs and its results,
and the offline demo. Figures are read from the result folders when the deck is built.

    python presentation/build_review3.py
"""
from __future__ import annotations

import json
import os
import statistics
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import build_review2 as r2  # noqa: E402
from build_review2 import (  # noqa: E402
    CODE_GREEN,
    CW,
    HEX,
    LM,
    ROOT,
    TOP,
    card,
    check_row,
    code,
    code_slide,
    column_chart,
    dark_card,
    hline,
    icon,
    notes_panel,
    num,
    pct,
    stat_card,
    table,
    text,
)
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from src.visualization.report import aggregate_metrics, load_results  # noqa: E402

OUTPUT = os.path.join(ROOT, "project_ppt_24bce2972_review3.pptx")
REVIEW1_DECK = "project_ppt_24bce2972 (1).pptx"
PENDING = "run in progress"
TOTAL_TESTS = 397
RECHECK_TOTAL = "539,170"       # executions in the independent re-check of the runs on real programs
CATEGORIES = r2.CATEGORIES

# The architecture slide: what each agent does now.
AGENTS = [
    ("1  Code Analysis Specialist", "code_analyzer.py", "the current TAC program",
     ["Control-flow graph builder",
      "Liveness analysis",
      "Available-copies analysis",
      "Six detectors: constant fold, algebraic, dead code, CSE, strength reduction, copy propagation",
      "Priority ranker"],
     "opportunities ranked by priority"),
    ("2  Optimization Specialist", "optimization_specialist.py", "the program and its opportunities",
     ["Router: rule, hybrid or llm mode",
      "Rule-based optimizer",
      "LLM optimizer: Gemini, Groq, Ollama",
      "Prompt builder and TAC reply parser",
      "Self-repair loop"],
     "one candidate program"),
    ("3  Verification Module", "verifier.py, smt_verifier.py", "current program + candidate",
     ["Input-set generator (six kinds of input sets)",
      "Reference TAC interpreter",
      "Differential tester",
      "Symbolic polynomial checker",
      "SMT prover (Z3)"],
     "PASS (tested or proved) or FAIL"),
    ("4  Cost Evaluator", "cost_evaluator.py", "program + verified candidate",
     ["Static counters and weight table",
      "Execution profiler: 25 fixed inputs",
      "Latency table per instruction",
      "Comparator: static or time cost model"],
     "cheaper, or no cost gain"),
]

CODE_REPAIR = '''
# orchestrator.py: a rejected LLM answer gets one more try
repairs_done = 0
while specialist.can_repair(proposal, repairs_done):
    repairs_done += 1
    logger.log_repair(iteration, proposal.problem)
    proposal = specialist.repair(current, proposal)
    outcome = self._judge(iteration, current, proposal, logger, proposals)

# optimizer_llm.py: what the model is told
REPAIR_TEMPLATE = """You proposed an optimization of the TAC below,
but your answer was rejected.

Original TAC:          {tac_string}
Your rejected answer:  {rejected}
Why it was rejected:   {problem}

Propose ONE semantics-preserving optimization of the ORIGINAL TAC
that does not have this problem. If you are not sure a change is
correct, make a smaller and safer one."""
'''

CODE_TIME = '''
LATENCY = {"copy": 1, "simple": 1, "multiply": 3, "divide": 20,
           "jump": 1, "branch": 2, "print": 1}        # cycles

def execution_time(self, program):
    latency = [instruction_latency(ins) for ins in program.instructions]
    sets = profile_inputs(sorted(program.get_input_variables()))
    total = 0.0
    for inputs in sets:                  # 20 random sets + 5 edge cases
        counts = [0] * len(latency)
        try:
            self.interpreter.execute(program, inputs, counts=counts)
        except TACRuntimeError:
            pass                         # a run that stops early still took time
        total += sum(c * cycles for c, cycles in zip(counts, latency))
    return total / len(sets)

def is_improvement(self, old, new):
    if self.model == "time" and new.exec_time != old.exec_time:
        return new.exec_time < old.exec_time
    return new.weighted_total < old.weighted_total
'''

CODE_SMT = '''
def smt_equivalent(original, candidate, timeout_ms=3000):
    runs_a = _outcomes(z3, original, input_var)    # every path of a program:
    runs_b = _outcomes(z3, candidate, input_var)   # (condition, how it ends, prints)

    def describe(runs):          # one (kind, count, values...) for the whole program
        fields = None
        for cond, kind, prints in reversed(runs):
            row = [IntVal(kind), IntVal(len(prints)), *prints, *padding]
            fields = row if fields is None else [
                If(cond, r, f) for r, f in zip(row, fields)]
        return fields

    solver = z3.Solver()
    solver.set("timeout", timeout_ms)
    solver.add(Or(*[x != y for x, y in zip(describe(runs_a), describe(runs_b))]))
    answer = solver.check()
    if answer == z3.unsat:
        return SmtResult("proved")           # no input tells them apart
    if answer == z3.sat:
        return SmtResult("refuted", counterexample=solver.model())
    return SmtResult("unknown")              # the tested verdict stands
'''

STRICT_C = '''
// Codeforces 4A, accepted submission
#include <stdio.h>
int main() {
    int w;
    scanf("%d", &w);
    if (w % 2 == 0 && w > 2) {
        puts("YES");
        return 0;
    }
    printf("NO");
    return 0;
}
'''

STRICT_TAC = '''
w = in_w
t0 = w % 2
t1 = t0 == 0
t2 = w > 2
t3 = t1 && t2
ifFalse t3 goto L0
print 1000000        # "YES"
goto L1              # return 0
L0:
print 1000001        # "NO"
L1:
'''

FINDING_C = '''
// Codeforces 219B, accepted submission
scanf("%I64d%I64d", &p, &d);
ans = p;
for (long long x = 10; x <= p; x *= 10) {
    long long a = p / x, b = p % x;
    if (b < x - 1) a--;
    if (p - d <= a * x + x - 1)
        ans = a * x + x - 1;
}
printf("%I64d\\n", ans);

// LLM: "constant folding: replaced the
// computation of x - 1 with its value 9"
//      t6 = b < 9          wrong once x is 100
'''

DEMO_BAT = '''
> ollama_demo.bat
  LLM: ollama qwen2.5-coder:7b - offline mode
  ...
  [21] LLM (Ollama)
       Constant folding: ... the value of `score`
       will always be 7.
       verified cost 9.0 → 3.5  ACCEPTED
  [22] COPY_PROPAGATION   score = 7 → uses of score
       verified cost 3.5 → 1.0  ACCEPTED
  ...
  cost 57.5 → 1.0   reduction 98.3%
  execution time: 24.0 → 1.0 cycles (95.8% less)
  equivalence: proved for all inputs
  verifier, final check against the original: PASS

> ollama_demo.bat my_program.c --inputs n=5
> ollama_demo.bat --llm-first
'''


# ---------------------------------------------------------------------- #
def load(folder: str) -> dict | None:
    """Metrics and per-program records of one run, or None if it has not finished."""
    try:
        records = load_results(os.path.join(ROOT, folder))
    except FileNotFoundError:
        return None
    return {"m": aggregate_metrics(records), "by_name": {r["program_name"]: r for r in records}}


def title_of(slide) -> str:
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            return shape.text_frame.text
    return ""


def replace_everywhere(deck, old: str, new: str) -> None:
    for slide in deck.order:
        for shape in slide.shapes:
            if shape.has_text_frame and old in shape.text_frame.text:
                for run in shape.text_frame._txBody.iter(qn("a:t")):
                    if run.text and old in run.text:
                        run.text = run.text.replace(old, new)


def mean(values) -> float:
    values = list(values)
    return statistics.mean(values) if values else 0.0


def value(run, getter, **opts):
    if run is None:
        return (PENDING, {"color": "a5"})
    result = getter(run)
    return (result, opts) if opts else result


def find_review1() -> str:
    for path in (os.path.join(ROOT, REVIEW1_DECK), os.path.join(os.path.dirname(ROOT), "Compiler", REVIEW1_DECK),
                 os.path.join(HERE, "Compiler_Optimization_Presentation.pptx")):
        if os.path.exists(path):
            return path
    raise FileNotFoundError("Review 1 deck not found")


def build() -> str:
    # ------------------------------------------------------------------ the Review 2 slides
    r2.AGENTS = AGENTS
    r2.find_base = find_review1
    held = {}
    real_save = r2.Deck.save
    r2.Deck.save = lambda self, path: held.setdefault("deck", self)
    try:
        r2.build()
    finally:
        r2.Deck.save = real_save
    deck = held["deck"]
    replace_everywhere(deck, "336", str(TOTAL_TESTS))
    replace_everywhere(deck, "loop controller (up to 50 iterations)", "loop controller (up to 50 iterations, 200 for real programs)")
    replace_everywhere(deck, "decision logger, final checker", "decision logger, final checker with revert, resume cache")
    tail = {"Project Status", "Conclusion", "Review 1 Comments"}
    references = deck.order[-1]
    deck.order = [s for s in deck.order[:-1] if title_of(s) not in tail]

    rule = load("output_real_rule_only")
    timed = load("output_real_rule_time")
    groq = load("output_real_groq")
    gemini = load("output_real_gemini")
    ollama = load("output_real_ollama")
    # the LLM runs that exist, in a fixed order (a run needs its provider's key or a local model)
    llm_runs = [entry for entry in (("Gemini", gemini, "a1", "gemini-3.1-flash-lite, Google cloud"),
                                    ("Groq", groq, "a4", "gpt-oss and qwen models, Groq cloud"),
                                    ("Ollama", ollama, "a2", "qwen2.5-coder:7b, this laptop")) if entry[1]]
    generated = load("output_rule_only")
    with open(os.path.join(ROOT, "dataset", "real_manifest.json"), encoding="utf-8") as fh:
        manifest = json.load(fh)
    as_c = sum(1 for e in manifest["programs"] if e["language"].startswith("GNU C") and "++" not in e["language"])

    def same_programs(run):
        """The rule-only average on exactly the programs an LLM run covered."""
        return mean(rule["by_name"][name]["reduction_pct"] for name in run["by_name"]) if run and rule else 0.0

    def improved(run):
        return sum(1 for name, rec in run["by_name"].items()
                   if rec["reduction_pct"] > rule["by_name"][name]["reduction_pct"] + 0.01) if run and rule else 0

    def consulted(run):
        """Programs on which the LLM was actually asked (its quota had not run out yet)."""
        return sum(1 for rec in run["by_name"].values() if rec.get("llm_calls", 0) > 0)

    def quota_note():
        short = [f"{name}'s free daily quota ended after {consulted(run)} programs; its other "
                 f"{run['m']['total_programs'] - consulted(run)} were optimized by the rules alone."
                 for name, run, _, _ in llm_runs if consulted(run) < 0.9 * run["m"]["total_programs"]]
        return " ".join(short)

    def llm_calls(run):
        return sum(rec.get("llm_calls", 0) for rec in run["by_name"].values())

    # ------------------------------------------------------------------ overview
    s = deck.new("Review 3: Four Additions to the System")
    additions = [
        ("redo", "a1", "LLM self-repair",
         "A rejected LLM answer is sent back with the verifier's counterexample or the syntax error, and the model "
         "gets one more try."),
        ("scale", "a4", "Execution time in the cost model",
         "The Cost Evaluator now also measures how long a program runs, so work inside a loop counts once per "
         "iteration."),
        ("shield", "a2", "SMT verification with Z3",
         "Programs with branches but no loops are proved equivalent for every input, not only tested on some."),
        ("layers", "a3", "A dataset of real programs",
         f"500 accepted Codeforces submissions from {manifest['problems']} problems, in the same seven categories "
         "as the generated set."),
    ]
    for i, (name, colour, heading, detail) in enumerate(additions):
        x, y = LM + (i % 2) * 6.19, TOP + (i // 2) * 2.05
        card(s, x, y, 5.94, 1.85)
        icon(s, name, x + 0.28, y + 0.3, d=0.62, fill=colour)
        text(s, x + 1.12, y + 0.24, 4.6, 0.42, heading, size=19, color="tx1", bold=True)
        text(s, x + 1.12, y + 0.72, 4.6, 1.05, detail, size=14)
    card(s, LM, TOP + 4.1, CW, 1.25, fill="a1", alpha=8)
    text(s, LM + 0.3, TOP + 4.1, CW - 0.6, 1.25, [[
        ("Also new: ", {"bold": True, "color": "a1"}),
        "runs that can be interrupted and resumed, a guard that keeps the original program if the end-to-end "
        "check ever fails, a one-command offline demo with the local model, and ",
        (f"{TOTAL_TESTS} automated tests", {"bold": True}), "."]], size=15, color="tx1", anchor="ctr")

    # ------------------------------------------------------------------ self-repair
    s = deck.new("LLM Self-Repair")
    steps = [("LLM proposes", "a1"), ("Verifier or parser rejects", "a3"), ("Reason is sent back", "a4"),
             ("LLM tries once more", "a1"), ("Same verify and cost gates", "a2")]
    bw, gap = 2.2, 0.283
    for i, (label, colour) in enumerate(steps):
        x = LM + i * (bw + gap)
        card(s, x, TOP, bw, 0.75, fill=colour, alpha=12, line=colour)
        text(s, x + 0.08, TOP, bw - 0.16, 0.75, label, size=14, color="tx1", bold=True, align="ctr", anchor="ctr")
        if i:
            hline(s, x - gap, TOP + 0.375, x, TOP + 0.375, color="a5", width=1.5, arrow=True)
    notes_panel(s, LM, TOP + 1.0, 6.0, 4.35, "What the model is told", [
        [("Wrong answer: ", {"bold": True}), "\"It changes what the program does: output mismatch on inputs "
                                             "{in_n: 5}: original printed [15], candidate printed [10].\""],
        [("Not valid TAC: ", {"bold": True}), "\"unrecognised instruction 'm = (in_n + 1) / 2'. One instruction "
                                              "has at most one operator.\""],
        [("No TAC at all: ", {"bold": True}), "\"The reply had no <tac>...</tac> block.\""],
        "A repair goes through exactly the same gates, so it cannot lower the safety of the system.",
    ], size=13.5, after=8)
    x = LM + 6.25

    def repairs(run):
        tried = run["m"].get("llm_repairs", 0)
        return f"{run['m'].get('llm_repairs_accepted', 0)} of {tried}" if tried else "0"

    width = (5.88 - 0.12 * (len(llm_runs) - 1)) / max(len(llm_runs), 1)
    for i, (name, run, colour, _) in enumerate(llm_runs):
        stat_card(s, x + i * (width + 0.12), TOP + 1.0, width, 2.05, repairs(run),
                  f"repairs accepted with {name}", colour=colour, big_size=24 if len(llm_runs) > 2 else 30)
    card(s, x, TOP + 3.3, 5.88, 2.05)
    text(s, x + 0.25, TOP + 3.45, 5.38, 1.8, [
        [("What it fixes. ", {"bold": True, "color": "a1"}),
         "Strong models mostly fail on form: they write a whole expression in one instruction. Told so, they "
         "answer again in valid TAC."],
        [("What it does not fix. ", {"bold": True, "color": "a3"}),
         "A small local model that got the logic wrong usually gets it wrong again, or plays safe and returns "
         "the program unchanged."]], size=13.5, color="tx1", after=8)

    code_slide(deck, "Code: The Self-Repair Loop", CODE_REPAIR, "How it works", [
        "The loop sits right after the gates: only an LLM answer that failed them can be repaired.",
        "The feedback is the verifier's own reason, with the inputs that show the mistake.",
        "One repair per rejected answer by default (--max-repairs).",
        "Repairs have their own limit and do not use up the call budget for fresh proposals.",
    ], size=11.5)

    # ------------------------------------------------------------------ execution time
    s = deck.new("Execution Time in the Cost Model")
    card(s, LM, TOP, 3.95, 5.35)
    text(s, LM + 0.25, TOP + 0.2, 3.5, 0.4, "Latency per executed instruction", size=17, color="a1", bold=True)
    table(s, LM + 0.25, TOP + 0.75, [2.35, 1.1], [
        ["Instruction", "Cycles"], ["copy, load a constant", "1"], ["+  -  <<  compare  logic", "1"],
        ["multiply", "3"], ["divide, modulo", "20"], ["goto", "1"], ["conditional branch", "2"], ["print", "1"],
        ["label", "0"]], size=13, row_h=0.42)
    text(s, LM + 0.25, TOP + 4.62, 3.5, 0.65, "Typical x86-64 figures, rounded. The same table is used before "
         "and after, so only the ratio matters.", size=12)
    notes_panel(s, LM + 4.2, TOP, 7.93, 2.95, "How the time is measured", [
        "The program is run on 25 fixed inputs: 20 random sets (half of them positive) and 5 edge cases.",
        "Each executed instruction adds its latency; the result is the average over the runs.",
        "An input's values depend only on its name, so the program and a candidate are timed on the same data.",
        "A loop body counts once per iteration, which the static cost cannot see.",
    ], size=13.5, after=6)
    card(s, LM + 4.2, TOP + 3.2, 7.93, 2.15, fill="a4", alpha=10)
    text(s, LM + 4.45, TOP + 3.35, 7.43, 1.9, [
        [("Two cost models. ", {"bold": True, "color": "a4"})],
        [("static (default): ", {"bold": True}), "a change is accepted if the weighted instruction count drops. "
                                                 "All Review 2 results use it and stay reproducible."],
        [("time: ", {"bold": True}), "a change is accepted if the measured time drops, or the time is equal and "
                                     "the static cost drops. Replacing x * 2 by x + x now counts as a gain."],
        "Every run reports both reductions, whichever model decides."], size=13.5, color="tx1", after=5)

    code_slide(deck, "Code: Measuring Execution Time", CODE_TIME, "How it works", [
        "The interpreter fills in how often each instruction ran.",
        "Time = executions x latency, summed over the program and averaged over the inputs.",
        "A run that traps is timed up to the trap.",
        "In the time model the static cost only breaks ties.",
    ], size=11.5)

    # ------------------------------------------------------------------ SMT
    s = deck.new("SMT Verification with Z3")
    stages = [("Follow every path", "A branch splits the path; a division splits off the case \"divisor is zero\"."),
              ("One formula per program", "(how the run ends, how many values it prints, the values) as a function "
                                          "of the inputs."),
              ("Ask Z3 for a difference", "Is there any input on which the two formulas differ?")]
    for i, (heading, detail) in enumerate(stages):
        x = LM + i * 4.127
        card(s, x, TOP, 3.88, 1.75)
        text(s, x + 0.25, TOP + 0.15, 3.4, 0.4, [[(f"{i + 1}  ", {"color": "a1"}), heading]], size=17, color="tx1", bold=True)
        text(s, x + 0.25, TOP + 0.65, 3.4, 1.0, detail, size=13.5)
        if i:
            hline(s, x - 0.247, TOP + 0.875, x, TOP + 0.875, color="a5", width=1.5, arrow=True)
    answers = [("unsat", "a2", "Proved", "No such input exists: the programs are equivalent for all inputs."),
               ("sat", "a3", "Refuted", "Z3 returns the input. It is replayed in the interpreter, then the change "
                                        "is rejected."),
               ("unknown", "a5", "No answer", "Timeout, usually non-linear arithmetic. The tested verdict stands.")]
    for i, (word, colour, heading, detail) in enumerate(answers):
        x = LM + i * 4.127
        card(s, x, TOP + 2.0, 3.88, 1.6, fill=colour, alpha=10)
        text(s, x + 0.25, TOP + 2.12, 3.4, 0.4, [[(word + "  ", {"font": "Consolas", "color": colour}), heading]],
             size=16, color="tx1", bold=True)
        text(s, x + 0.25, TOP + 2.6, 3.4, 0.95, detail, size=13.5)
    card(s, LM, TOP + 3.85, CW, 1.5)
    text(s, LM + 0.3, TOP + 3.95, CW - 0.6, 1.3, [
        [("What testing misses. ", {"bold": True, "color": "a1"}),
         "A program prints 1 when 3x = 999999 and 0 otherwise; the candidate always prints 0. They differ on "
         "one input, x = 333333, which is near no literal in either program, so every test passes. Z3 finds it."],
        [("Where it runs. ", {"bold": True, "color": "a1"}),
         "On every LLM proposal and on the end-to-end check of every program. Loops are outside its scope and "
         "stay tested. Z3 is optional: without it the stage is skipped."]], size=13.5, color="tx1", after=7)

    code_slide(deck, "Code: The SMT Equivalence Check", CODE_SMT, "How it works", [
        "_outcomes walks the program symbolically and returns every way a run can end.",
        "The path conditions are exclusive and cover everything, so nested If() describes the whole program.",
        "Integers are unbounded and / and % follow C, exactly as in the interpreter.",
        "More than 256 paths, or a loop: unsupported, and testing decides.",
    ], size=11.5)

    # ------------------------------------------------------------------ real dataset
    s = deck.new("A Dataset of Real Programs")
    funnel = [("12.6 M", "submissions in open-r1/codeforces-submissions (Hugging Face, 2025, CC-BY-4.0)", "a5"),
              ("469,491", "accepted C and C++ submissions examined, from 6 of the 36 data files", "a1"),
              ("9,489", "usable: compile in strict mode, print something, finish on every test input", "a4"),
              ("500", f"chosen, from {manifest['problems']} problems, newest submission first", "a2")]
    for i, (big, label, colour) in enumerate(funnel):
        y = TOP + i * 1.34
        card(s, LM, y, 5.6, 1.16, fill=colour, alpha=10)
        text(s, LM + 0.25, y, 1.85, 1.16, big, size=26, color=colour, bold=True, anchor="ctr")
        text(s, LM + 2.15, y, 3.3, 1.16, label, size=13, color="tx1", anchor="ctr")
    rows = [["Category", "Programs", "Problems", "Decided by"]]
    decided = {"arithmetic": "straight-line, shallow", "nested": "arithmetic 3+ operators deep",
               "repeated_expr": "an expression written twice", "dead_code": "a stored value never read",
               "loops": "loop, no branch", "conditionals": "branch, no loop", "mixed": "loop and branch"}
    for key, label in CATEGORIES:
        entries = [e for e in manifest["programs"] if e["category"] == key]
        rows.append([label, str(len(entries)), str(len({e["problem_id"] for e in entries})), decided[key]])
    table(s, LM + 5.85, TOP, [1.75, 1.05, 1.05, 2.43], rows, size=13, row_h=0.44)
    text(s, LM + 5.85, TOP + 3.72, 6.28, 1.65, [
        f"The proportions are the generator's: 60 / 60 / 60 / 60 / 80 / 80 / 100.",
        f"{as_c} were submitted as C and {500 - as_c} as C-style C++ (scanf / printf, nothing else from C++).",
        "Programs that compile to the same TAC count once; each category is filled one program per problem "
        "per round.",
        "Only 2 % qualify: arrays, other functions and input read inside a loop are not supported.",
    ], size=13, color="tx1", after=5, bullets=True)

    s = deck.new("Strict Mode: Real C to TAC")
    code(s, LM, TOP, 4.5, 4.1, STRICT_C, size=12.5, lang="c")
    code(s, LM + 4.75, TOP, 3.55, 4.1, STRICT_TAC, size=12.5, lang="py")
    hline(s, LM + 4.5, TOP + 2.05, LM + 4.75, TOP + 2.05, color="a5", width=1.5, arrow=True)
    notes_panel(s, LM + 8.55, TOP, 3.58, 5.35, "What strict mode keeps", [
        "Printed text is output: each distinct text gets a code, so changing which message is printed fails "
        "verification.",
        "return in the middle of main and exit() stop the program.",
        "Also read now: character literals, bool, long long aliases, hex and octal numbers.",
        "a && b is refused when b divides: TAC has no short-circuit.",
        "A macro other than a type alias is refused.",
    ], size=13, after=7)
    card(s, LM, TOP + 4.3, 8.3, 1.05, fill="a1", alpha=8)
    text(s, LM + 0.25, TOP + 4.3, 7.8, 1.05, "The demo does not use strict mode, so its output for the five "
         "Review 2 sample programs is unchanged.", size=13.5, color="tx1", anchor="ctr")

    # ------------------------------------------------------------------ results: key metrics
    s = deck.new("Real Programs: Key Metrics")
    columns = [("Rules, static", rule), ("Rules, time", timed)] + [(f"Rules + {name}", run) for name, run, _, _ in llm_runs]

    def row(label, getter, llm_only=False, **opts):
        cells = [label]
        for name, run in columns:
            if llm_only and "+" not in name:
                cells.append("n/a")
            else:
                cells.append(value(run, getter, **opts))
        return cells

    rows = [["Metric"] + [name for name, _ in columns]]
    rows.append(row("Programs", lambda r: num(r["m"]["total_programs"])))
    rows.append(row("Programs the LLM was consulted on", lambda r: num(consulted(r)), llm_only=True))
    rows.append(row("Average cost reduction", lambda r: pct(r["m"]["avg_cost_reduction_pct"]), bold=True, color="a1"))
    rows.append(row("Rules alone on the same programs", lambda r: pct(same_programs(r))))
    rows.append(row("Average execution-time reduction", lambda r: pct(r["m"]["avg_time_reduction_pct"])))
    rows.append(row("Proved equivalent for all inputs", lambda r: f"{r['m']['programs_proved_equivalent']} "
                                                                  f"({r['m']['programs_proved_by_smt']} by SMT)"))
    rows.append(row("Total proposals", lambda r: num(r["m"]["total_proposals"])))
    rows.append(row("LLM calls", lambda r: num(llm_calls(r)), llm_only=True))
    rows.append(row("LLM validity rate", lambda r: pct(r["m"]["llm_validity_rate_pct"]), llm_only=True))
    rows.append(row("Self-repairs accepted / tried", lambda r: f"{r['m'].get('llm_repairs_accepted', 0)} / "
                                                               f"{r['m'].get('llm_repairs', 0)}", llm_only=True))
    rows.append(row("Wrong LLM proposals blocked", lambda r: num(r["m"]["by_source"]["LLM"]["invalid"]), llm_only=True))
    rows.append(row("False positives", lambda r: str(r["m"]["false_positive_count"]), bold=True, color="a2"))
    first = 3.38 if len(columns) > 4 else 3.93
    table(s, LM, TOP, [first] + [(CW - first) / len(columns)] * len(columns), rows, size=12.5 if len(columns) > 4 else 13,
          row_h=0.365)
    text(s, LM, TOP + 4.9, CW, 0.45,
         "LLM budget per program: 2 fresh proposals for every fifth program and 1 for the others, plus 1 repair per "
         "rejected answer (this keeps the cloud runs inside their free daily quotas). " + quota_note(), size=12)

    # ------------------------------------------------------------------ results: by category
    s = deck.new("Real Programs: Cost Reduction by Category")
    series = []
    if generated:
        series.append(("Generated dataset, rules", [generated["m"]["per_category"][k]["avg_reduction_pct"] for k, _ in CATEGORIES], HEX["a5"]))
    if rule:
        series.append(("Real programs, rules", [rule["m"]["per_category"][k]["avg_reduction_pct"] for k, _ in CATEGORIES], HEX["a1"]))
    for name, run, colour in ((f"Real programs, rules + {name}", run, HEX[colour if colour != "a1" else "a3"])
                              for name, run, colour, _ in llm_runs):
        if run:
            # a category the run has not finished is left out (no bar) instead of drawn from a few programs
            done = {k: sum(1 for r in run["by_name"].values() if r["category"] == k) for k, _ in CATEGORIES}
            full = {k: rule["m"]["per_category"][k]["programs"] for k, _ in CATEGORIES}
            series.append((name, [run["m"]["per_category"][k]["avg_reduction_pct"] if done[k] == full[k] else None
                                  for k, _ in CATEGORIES], colour))
    column_chart(s, LM, TOP, 8.85, 5.3, [label for _, label in CATEGORIES], series)
    r2.what_it_shows(s, [
        "Generated programs fold to constants; real ones read their input, so less can be removed.",
        f"Real programs, rules: {pct(rule['m']['avg_cost_reduction_pct']) if rule else PENDING} on average.",
        *([f"{llm_runs[0][0]} with self-repair improves {improved(llm_runs[0][1])} of the "
           f"{llm_runs[0][1]['m']['total_programs']} real programs: real loops depend on the input, so few can be "
           "folded away."] if llm_runs else []),
    ])

    # ------------------------------------------------------------------ results: static vs time
    s = deck.new("Real Programs: Static vs. Time Cost Model")

    def kind(run, name, field="accepted"):
        return run["m"]["by_optimization_type"].get(name, {}).get(field, 0)

    rows = [["", "Static cost model", "Time cost model"]]
    rows.append(["Average static cost reduction", value(rule, lambda r: pct(r["m"]["avg_cost_reduction_pct"])),
                 value(timed, lambda r: pct(r["m"]["avg_cost_reduction_pct"]))])
    rows.append(["Average execution-time reduction", value(rule, lambda r: pct(r["m"]["avg_time_reduction_pct"])),
                 value(timed, lambda r: pct(r["m"]["avg_time_reduction_pct"]), bold=True, color="a1")])
    rows.append(["Median execution-time reduction", value(rule, lambda r: pct(r["m"]["median_time_reduction_pct"])),
                 value(timed, lambda r: pct(r["m"]["median_time_reduction_pct"]))])
    rows.append(["Strength reductions accepted", value(rule, lambda r: num(kind(r, "STRENGTH_REDUCTION"))),
                 value(timed, lambda r: num(kind(r, "STRENGTH_REDUCTION")))])
    rows.append(["Strength reductions rejected (no gain)", value(rule, lambda r: num(kind(r, "STRENGTH_REDUCTION", "rejected"))),
                 value(timed, lambda r: num(kind(r, "STRENGTH_REDUCTION", "rejected")))])
    rows.append(["Proposals accepted", value(rule, lambda r: num(r["m"]["accepted"])), value(timed, lambda r: num(r["m"]["accepted"]))])
    rows.append(["False positives", value(rule, lambda r: str(r["m"]["false_positive_count"]), bold=True, color="a2"),
                 value(timed, lambda r: str(r["m"]["false_positive_count"]), bold=True, color="a2")])
    table(s, LM, TOP, [4.3, 2.3, 2.3], rows, size=14, row_h=0.5)
    notes_panel(s, LM + 9.15, TOP, 2.98, 5.35, "Reading it", [
        "Same rules, same 500 programs; only the definition of \"cheaper\" changes.",
        "The time model accepts changes the static count calls a tie, such as x * 2 to x + x.",
        "Time falls less than size on real programs: most removed instructions are copies, which cost one "
        "cycle each.",
    ], size=13, after=8)
    text(s, LM, TOP + 4.25, 8.9, 1.0, "Execution time is measured on 25 inputs per program with the latency "
         "table of the cost evaluator; both runs allow 200 iterations per program.", size=12.5)

    # ------------------------------------------------------------------ results: proved
    s = deck.new("Real Programs: What Was Proved")
    proved = rule["m"]["programs_proved_equivalent"] if rule else 0
    by_smt = rule["m"]["programs_proved_by_smt"] if rule else 0
    stat_card(s, LM, TOP, 3.88, 2.1, f"{proved} / 500" if rule else PENDING,
              "optimized programs proved equal to their original for every input", colour="a2", big_size=34)
    stat_card(s, LM, TOP + 2.3, 3.88, 1.45, str(by_smt) if rule else PENDING, "of them by the SMT solver", colour="a1", big_size=30)
    stat_card(s, LM, TOP + 3.95, 3.88, 1.4, "0", "programs reverted by the final check", colour="a4", big_size=30)
    rows = [["Category", "Programs", "Proved", "How"]]
    how = {"arithmetic": "polynomial check", "nested": "polynomial check", "repeated_expr": "polynomial or SMT",
           "dead_code": "polynomial or SMT", "loops": "tested (loops)", "conditionals": "SMT", "mixed": "tested (loops)"}
    for key, label in CATEGORIES:
        cat = rule["m"]["per_category"][key] if rule else None
        rows.append([label, str(cat["programs"]) if cat else "", str(cat.get("programs_proved", 0)) if cat else PENDING, how[key]])
    table(s, LM + 4.13, TOP, [2.1, 1.25, 1.15, 3.5], rows, size=13.5, row_h=0.46)
    text(s, LM + 4.13, TOP + 3.85, 8.0, 1.5, [
        "A program with a loop cannot be proved by this method, so it stays tested on hundreds of inputs.",
        "For programs without loops the claim is stronger than zero false positives on the tests: no input "
        "exists on which the optimized program behaves differently.",
        "An independent re-check of all 500 pairs on other, wider inputs: 107,834 executions, 0 mismatches.",
    ], size=13, color="tx1", after=6, bullets=True)

    # ------------------------------------------------------------------ results: LLMs
    s = deck.new("Real Programs: The LLMs with Self-Repair")
    wide = len(llm_runs) < 3
    cw = (CW - 0.25 * (len(llm_runs) - 1)) / max(len(llm_runs), 1)
    for i, (name, run, colour, where) in enumerate(llm_runs):
        x = LM + i * (cw + 0.25)
        card(s, x, TOP, cw, 5.35)
        text(s, x + 0.22, TOP + 0.2, cw - 0.44, 0.42, "Rules + " + name, size=19, color=colour, bold=True)
        text(s, x + 0.22, TOP + 0.66, cw - 0.44, 0.3, where, size=11.5)
        m = run["m"]
        base = same_programs(run)
        lines = [
            ["Programs", num(m["total_programs"])],
            ["LLM consulted on", num(consulted(run))],
            ["Rules alone on them", pct(base)],
            ["With the LLM", (pct(m["avg_cost_reduction_pct"]), {"bold": True, "color": colour})],
            ["Programs improved", f"{improved(run)}"],
            ["LLM calls", num(llm_calls(run))],
            ["Accepted / no gain / wrong", f"{m['by_source']['LLM']['accepted']} / {m['by_source']['LLM']['rejected']} / "
                                            f"{m['by_source']['LLM']['invalid']}"],
            ["Validity rate", pct(m["llm_validity_rate_pct"])],
            ["Repairs accepted / tried", f"{m.get('llm_repairs_accepted', 0)} / {m.get('llm_repairs', 0)}"],
            ["False positives", (str(m["false_positive_count"]), {"bold": True, "color": "a2"})],
        ]
        table(s, x + 0.22, TOP + 1.1, [(cw - 0.44) * 0.65, (cw - 0.44) * 0.35], lines, size=13 if wide else 12,
              row_h=0.4, header=False)

    # ------------------------------------------------------------------ the wrong change the re-check found
    s = deck.new("Validation: A Wrong Change the Tests Missed")
    code(s, LM, TOP, 5.3, 5.35, FINDING_C, size=12, lang="c")
    steps = [("a3", "What happened", "In the Groq run the model replaced x - 1 by 9. That is right while x is 10, and "
                                     "the loop runs a second time only for p of 100 or more. No input set of the "
                                     "verifier reached that, so all 388 tests passed and the change was accepted."),
             ("a4", "How it was found", "Every finished run is re-checked on inputs the verifier does not use, with "
                                        "values up to a million. One program out of 100 disagreed: for p = 779, "
                                        "d = 428 the original prints 699 and the optimized program 799."),
             ("a2", "The fix", "The verifier now adds 60 wide-range input sets: values up to a thousand, ten "
                               "thousand, a million and a billion, signed values, and powers of ten with their "
                               "neighbours. The wrong change is now rejected with a counterexample.")]
    for i, (colour, heading, detail) in enumerate(steps):
        y = TOP + i * 1.5
        card(s, LM + 5.55, y, 6.58, 1.36, fill=colour, alpha=10)
        text(s, LM + 5.8, y + 0.1, 6.1, 0.34, heading, size=15, color=colour, bold=True)
        text(s, LM + 5.8, y + 0.46, 6.1, 0.9, detail, size=12.5, color="tx1")
    card(s, LM + 5.55, TOP + 4.5, 6.58, 0.85)
    text(s, LM + 5.8, TOP + 4.5, 6.1, 0.85, [[("After the fix: ", {"bold": True, "color": "a1"}),
                                             f"the program was optimized again, and the re-check of every run on real "
                                             f"programs ({RECHECK_TOTAL} executions) finds 0 mismatches. The four "
                                             "generated-dataset runs were re-checked the same way: 0 mismatches."]],
         size=12.5, color="tx1", anchor="ctr")

    # ------------------------------------------------------------------ offline demo
    s = deck.new("Offline Demo with the Local Model")
    code(s, LM, TOP, 6.6, 5.35, DEMO_BAT, size=12, lang="log")
    notes_panel(s, LM + 6.85, TOP, 5.28, 5.35, "One command, no internet", [
        "ollama_demo.bat checks that Ollama is running and the model is pulled, then starts the demo in "
        "offline mode.",
        "Without a file it runs sample 6: a branch that is always taken, which the rules leave behind and the "
        "7B model can usually fold.",
        "When the model is wrong, the verifier blocks it on screen and self-repair gives it a second try.",
        "With the updated NVIDIA driver the model runs mostly on the GPU, about 20 tokens per second.",
        "--llm-first asks the model before the rules, for more LLM steps to watch.",
    ], size=13, after=7)

    # ------------------------------------------------------------------ status and conclusion
    s = deck.new("Project Status")
    status = [("check", "a2", "System built", ["Front end: C (incl. strict mode) and mini-language to TAC",
                                              "Orchestrator loop + decision logging", "Six rule-based optimizations",
                                              "Router with three LLM backends", "LLM self-repair",
                                              "Offline mode and one-command local demo", "Resumable full runs"]),
              ("shield", "a1", "Verified and measured", ["Differential + symbolic + SMT verification",
                                                        f"{TOTAL_TESTS} automated tests, all passing",
                                                        "5,100 random programs fuzzed, 800 mutation tests",
                                                        "500 generated programs, 4 full runs",
                                                        "500 real programs, rule runs under 2 cost models",
                                                        "LLM runs with self-repair on the real programs",
                                                        "Static cost and execution time reported"]),
              ("route", "a4", "Limits and next steps", ["Arrays and functions in the front end",
                                                       "SMT for loops (bounded unrolling)",
                                                       "32-bit overflow is not modelled",
                                                       "A larger LLM budget per real program",
                                                       "Rule passes for loops and branches"])]
    for i, (name, colour, heading, items) in enumerate(status):
        x = LM + i * 4.127
        card(s, x, TOP, 3.88, 5.35)
        icon(s, name, x + 0.25, TOP + 0.22, d=0.5, fill=colour)
        text(s, x + 0.9, TOP + 0.27, 2.8, 0.42, heading, size=18, color="tx1", bold=True)
        text(s, x + 0.25, TOP + 0.95, 3.4, 4.3, items, size=14, color="tx1", after=9, bullets=True)

    s = deck.new("Conclusion", kind="DARK")

    def avg(run):
        return pct(run["m"]["avg_cost_reduction_pct"], 1) if run else "pending"

    conclusion = [
        (avg(generated), "average cost reduction on the 500 generated programs, rules alone", "a6"),
        (avg(rule), "on 500 real Codeforces programs, rules alone", "a6"),
        (pct(timed["m"]["avg_time_reduction_pct"], 1) if timed else "pending",
         "less execution time on the real programs under the time cost model", "a6"),
        (f"{proved} / 500" if rule else "pending", "real programs proved equivalent for every input, not only tested", "a6"),
        (avg(max((r for _, r, _, _ in llm_runs), key=lambda r: r["m"]["avg_cost_reduction_pct"], default=None)),
         "best LLM configuration on the real programs, with self-repair"
         + (f" ({pct(same_programs(llm_runs[0][1]), 1)} with rules alone)" if llm_runs else ""), "a6"),
        ("1 → 0", "wrong change that passed the tests on real programs: caught by the re-check, now rejected",
         CODE_GREEN),
    ]
    for i, (big, label, colour) in enumerate(conclusion):
        dark_card(s, LM + (i % 3) * 4.127, TOP + (i // 3) * 2.2, 3.88, 1.95, big, label, colour)
    text(s, LM, 5.9, CW, 0.75, [[("Key insight: ", {"bold": True, "color": "a6"}),
                                 "any proposer can be used, rule, cloud model or laptop model, because nothing "
                                 "is accepted until it is tested, and where possible proved."]], size=18, color="bg1")

    # references: two more
    column = next(sh for sh in references.shapes if sh.has_text_frame and sh.text_frame.text.startswith("[7]"))
    last = column.text_frame.paragraphs[-1]._p
    import copy
    for entry in ("[15]  de Moura, L., Bjorner, N. (2008). Z3: An Efficient SMT Solver. TACAS 2008.",
                  "[16]  open-r1 (2025). codeforces-submissions (dataset, CC-BY-4.0). Hugging Face."):
        clone = copy.deepcopy(last)
        runs = clone.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = entry
        for extra in runs[1:]:
            clone.remove(extra)
        last.addnext(clone)
        last = clone
    deck.order.append(references)

    try:
        deck.save(OUTPUT)
        return OUTPUT
    except PermissionError:
        fallback = OUTPUT.replace(".pptx", "_new.pptx")
        deck.save(fallback)
        return fallback + "  (the main file is open in PowerPoint; close it and rebuild to replace it)"


if __name__ == "__main__":
    print("saved", build())
