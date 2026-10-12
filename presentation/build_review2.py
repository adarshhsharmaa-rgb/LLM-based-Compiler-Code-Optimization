"""Build the Review 2 deck from the Review 1 deck.

    py presentation/build_review2.py

Slide 1 and the theme, layouts and icons come from the Review 1 file
(``project_ppt_24bce2972*.pptx`` in the project root, or the copy kept in this
folder). Numbers are read from the result folders at build time, so slides
about a run that has not finished yet (Groq) say so and fill in on a rebuild.
"""
from __future__ import annotations

import copy
import glob
import hashlib
import io
import os
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from src.visualization.report import aggregate_metrics, compare_runs, load_results  # noqa: E402

OUTPUT = os.path.join(ROOT, "project_ppt_24bce2972_review2.pptx")
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# ---------------------------------------------------------------------- #
# theme (same scheme colours as the Review 1 deck)
# ---------------------------------------------------------------------- #
SCHEME = {
    "tx1": MSO_THEME_COLOR.TEXT_1, "tx2": MSO_THEME_COLOR.TEXT_2,
    "bg1": MSO_THEME_COLOR.BACKGROUND_1, "bg2": MSO_THEME_COLOR.BACKGROUND_2,
    "a1": MSO_THEME_COLOR.ACCENT_1, "a2": MSO_THEME_COLOR.ACCENT_2, "a3": MSO_THEME_COLOR.ACCENT_3,
    "a4": MSO_THEME_COLOR.ACCENT_4, "a5": MSO_THEME_COLOR.ACCENT_5, "a6": MSO_THEME_COLOR.ACCENT_6,
}
HEX = {"tx1": "1B2A41", "tx2": "4A5568", "bg2": "EEF3FA", "a1": "2A78D6", "a2": "0B8A0B",
       "a3": "C23434", "a4": "C98500", "a5": "6B7A90", "a6": "BFD6F6"}
CODE_BG, CODE_FG, CODE_DIM, CODE_BLUE, CODE_GREEN = "111827", "E5E7EB", "9CA3AF", "8FB8F0", "7FD27F"
CODE_RED, CODE_AMBER = "F28B82", "F5C06B"
BORDER = "D5DCE6"
ICONS = {"arrow": "c13de758", "check": "0f58fd08", "bulb": "0440cd4a", "shield": "f28c5a1c",
         "code": "ccc4ce4d", "search": "2b698dbc", "layers": "0e1f4077", "redo": "fb931792",
         "list": "6d10d92b", "chat": "c1fc2bd6", "flag": "957d4dfc", "wrench": "cfda1c49",
         "scale": "b7931ad1", "chart": "707fae7b", "route": "cec23f96"}
LM, CW, TOP = 0.6, 12.13, 1.4          # left margin, content width, content top (inches)


def E(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


def paint(color_format, color: str) -> None:
    if color in SCHEME:
        color_format.theme_color = SCHEME[color]
    else:
        color_format.rgb = RGBColor.from_string(color)


# ---------------------------------------------------------------------- #
# drawing helpers
# ---------------------------------------------------------------------- #
def _plain(shape):
    """Drop the theme style reference so no default outline, shadow or font colour sneaks in."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def card(slide, x, y, w, h, fill="bg2", alpha=None, radius=0.08, line=None, line_w=1.25, shape=None):
    shp = _plain(slide.shapes.add_shape(shape or MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h)))
    if shape is None:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.fill.solid()
    paint(shp.fill.fore_color, fill)
    if alpha is not None:
        clr = shp._element.spPr.find(qn("a:solidFill"))[0]
        a = OxmlElement("a:alpha")
        a.set("val", str(int(alpha * 1000)))
        clr.append(a)
    if line:
        paint(shp.line.color, line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    return shp


def text(slide, x, y, w, h, content, size=14, color="tx2", bold=False, font=None, align=None,
         anchor="t", after=None, bullets=False, line=None):
    """content: a paragraph or list of paragraphs; a paragraph is a string or a list of runs;
    a run is a string or (string, {size, bold, color, font, italic})."""
    box = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for run in (para if isinstance(para, list) else [para]):
            run_text, opts = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = run_text
            r.font.size = Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", bold)
            if opts.get("italic"):
                r.font.italic = True
            paint(r.font.color, opts.get("color", color))
            face = opts.get("font", font)
            if face:
                r.font.name = face
            r._r.find(qn("a:t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if align:
            p.alignment = {"ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "l": PP_ALIGN.LEFT}[align]
        if after is not None:
            p.space_after = Pt(after)
        if line:
            p.line_spacing = line
        if bullets:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "228600")
            ppr.set("indent", "-228600")
            bu = OxmlElement("a:buChar")
            bu.set("char", "•")
            ppr.append(bu)
    return box


def picture(slide, source, x, y, w, h):
    """Place an image inside the box (x, y, w, h), keeping its proportions, centred."""
    stream = io.BytesIO(source) if isinstance(source, bytes) else source
    with Image.open(stream) as img:
        ratio = img.width / img.height
    if isinstance(stream, io.BytesIO):
        stream.seek(0)
    pw, ph = (w, w / ratio) if w / ratio <= h else (h * ratio, h)
    return slide.shapes.add_picture(stream, E(x + (w - pw) / 2), E(y + (h - ph) / 2), E(pw), E(ph))


def icon(slide, name, x, y, d=0.55, fill="a1"):
    """An icon from the Review 1 deck in a filled circle with its top-left corner at (x, y)."""
    card(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    g = d * 0.52
    picture(slide, ICON_BLOBS[ICONS[name]], x + (d - g) / 2, y + (d - g) / 2, g, g)


def hline(slide, x1, y1, x2, y2, color="a5", width=1.0, arrow=False):
    """A straight line; with ``arrow`` it ends in an arrowhead at (x2, y2)."""
    ln = _plain(slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)))
    paint(ln.line.color, color)
    ln.line.width = Pt(width)
    if arrow:
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        ln.line._get_or_add_ln().append(tail)
    return ln


def box(slide, x, y, w, h, label, fill="bg2", line="a1", size=13, bold=True, color="tx1", radius=0.07):
    """A small outlined box with centred text (pipeline stages in the architecture diagram)."""
    card(slide, x, y, w, h, fill=fill, line=line, line_w=1.25, radius=radius)
    text(slide, x + 0.04, y, w - 0.08, h, label, size=size, color=color, bold=bold, align="ctr", anchor="ctr")


def _border(tag):
    return parse_xml(
        f'<a:{tag} xmlns:a="{A_NS}" w="9525" cap="flat" cmpd="sng" algn="ctr"><a:solidFill>'
        f'<a:srgbClr val="{BORDER}"/></a:solidFill><a:prstDash val="solid"/><a:round/>'
        f'<a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/></a:{tag}>')


def table(slide, x, y, widths, rows, size=14, row_h=0.47, header=True, first_col_bold=False):
    """rows: lists of cells; a cell is a string or (string, {bold, color, align})."""
    frame = slide.shapes.add_table(len(rows), len(widths), E(x), E(y), E(sum(widths)), E(row_h * len(rows)))
    tbl = frame.table
    tbl_pr = tbl._tbl.tblPr
    for key in list(tbl_pr.attrib):
        del tbl_pr.attrib[key]
    for child in list(tbl_pr):
        tbl_pr.remove(child)
    for c, width in enumerate(widths):
        tbl.columns[c].width = E(width)
    for r, row in enumerate(rows):
        tbl.rows[r].height = E(row_h)
        for c, value in enumerate(row):
            cell_text, opts = (value, {}) if isinstance(value, str) else value
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(88900)
            cell.margin_top = cell.margin_bottom = Emu(50800)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = cell_text
            run.font.size = Pt(size)
            run.font.name = "Calibri"
            run.font.bold = opts.get("bold", (header and r == 0) or (first_col_bold and c == 0))
            paint(run.font.color, opts.get("color", "tx1"))
            if opts.get("align"):
                p.alignment = {"ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[opts["align"]]
            tc_pr = cell._tc.get_or_add_tcPr()
            for child in list(tc_pr):
                tc_pr.remove(child)
            for tag in ("lnL", "lnR", "lnT", "lnB"):
                tc_pr.append(_border(tag))
            fill = opts.get("fill", "bg2" if header and r == 0 else None)
            if fill:
                tc_pr.append(parse_xml(f'<a:solidFill xmlns:a="{A_NS}"><a:schemeClr val="{fill}"/></a:solidFill>'))
            else:
                tc_pr.append(parse_xml(f'<a:noFill xmlns:a="{A_NS}"/>'))
    return frame


# ---------------------------------------------------------------------- #
# code and terminal panels
# ---------------------------------------------------------------------- #
_PY = re.compile(
    r"(?P<comment>#.*$)|(?P<string>[fr]?\"(?:[^\"\\]|\\.)*\"|[fr]?'(?:[^'\\]|\\.)*')|"
    r"(?P<kw>\b(?:def|class|return|if|elif|else|for|while|in|not|and|or|is|None|True|False|break|"
    r"continue|import|from|raise|try|except|with|as|lambda)\b)")
_C = re.compile(r"(?P<comment>/\*.*?\*/|//.*$)|(?P<string>\"(?:[^\"\\]|\\.)*\"|<[a-z.]+>)|"
                r"(?P<kw>\b(?:int|for|while|if|else|return)\b|#include)")
_LOG_TAGS = {"[ACCEPT]": CODE_GREEN, "[REJECT]": CODE_RED, "[ROUTE]": CODE_AMBER, "[LLM]": CODE_BLUE,
             "[INFO]": CODE_DIM}


def _tokens(line, pattern):
    runs, pos = [], 0
    for m in pattern.finditer(line):
        if m.start() > pos:
            runs.append((line[pos:m.start()], CODE_FG, False))
        kind = m.lastgroup
        colour = {"comment": CODE_DIM, "string": CODE_GREEN, "kw": CODE_BLUE}[kind]
        runs.append((m.group(0), colour, False))
        pos = m.end()
    if pos < len(line):
        runs.append((line[pos:], CODE_FG, False))
    return runs or [(" ", CODE_FG, False)]


def _log_runs(line):
    stripped = line.strip()
    if stripped.startswith(("===", "STAGE")):
        return [(line, CODE_BLUE, True)]
    if stripped.startswith("---") or stripped.startswith("..."):
        return [(line, CODE_DIM, False)]
    for tag, colour in _LOG_TAGS.items():
        if stripped.startswith(tag):
            head = line[:line.index(tag) + len(tag)]
            return [(head, colour, True), (line[len(head):], CODE_GREEN if tag == "[ACCEPT]" else CODE_FG, False)]
    runs, pos = [], 0
    for m in re.finditer(r"ACCEPTED|PASS\b|identical output|REJECTED[^\n]*|FAILED|FAIL\b", line):
        if m.start() > pos:
            runs.append((line[pos:m.start()], CODE_FG, False))
        bad = m.group(0).startswith(("REJECT", "FAIL"))
        runs.append((m.group(0), CODE_RED if bad else CODE_GREEN, True))
        pos = m.end()
    if pos < len(line):
        runs.append((line[pos:], CODE_FG, False))
    if stripped.startswith(("route to", "Found", "No more", "LLM not consulted", "Routing:", "LLM:")):
        runs = [(t, CODE_DIM if c == CODE_FG else c, b) for t, c, b in runs]
    return runs or [(" ", CODE_FG, False)]


def code(slide, x, y, w, h, source, size=12, lang="py", pad=0.22):
    """A dark rounded panel with monospaced, lightly highlighted text."""
    card(slide, x, y, w, h, fill=CODE_BG, radius=0.08)
    paragraphs = []
    for line in source.strip("\n").split("\n"):
        if lang == "py":
            runs = _tokens(line, _PY)
        elif lang == "c":
            runs = _tokens(line, _C)
        elif lang == "log":
            runs = _log_runs(line)
        else:
            runs = [(line or " ", CODE_FG, False)]
        paragraphs.append([(t, {"color": c, "bold": b}) for t, c, b in runs])
    return text(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, paragraphs, size=size, color=CODE_FG,
                font="Consolas", after=1)


# ---------------------------------------------------------------------- #
# deck plumbing
# ---------------------------------------------------------------------- #
def find_base() -> str:
    candidates = sorted(glob.glob(os.path.join(ROOT, "project_ppt_24bce2972*.pptx")))
    candidates = [c for c in candidates if "review2" not in os.path.basename(c).lower()]
    candidates.append(os.path.join(HERE, "Compiler_Optimization_Presentation.pptx"))
    for path in candidates:
        if os.path.exists(path):
            return path
    raise FileNotFoundError("Review 1 deck not found")


def placeholder(slide, kind):
    return next(sh for sh in slide.shapes if sh.is_placeholder and sh.placeholder_format.type == kind)


class Deck:
    def __init__(self, base: str):
        self.prs = Presentation(base)
        self.original = list(self.prs.slides)
        self.proto = {}
        for name, index in (("CONTENT", 3), ("DARK", 23)):
            src = self.original[index]
            self.proto[name] = (copy.deepcopy(placeholder(src, PP_PLACEHOLDER.TITLE)._element),
                                copy.deepcopy(placeholder(src, PP_PLACEHOLDER.SLIDE_NUMBER)._element))
        self.order = []

    def keep(self, number: int, title: str | None = None):
        slide = self.original[number - 1]
        if title:
            placeholder(slide, PP_PLACEHOLDER.TITLE).text_frame.paragraphs[0].runs[0].text = title
        self.order.append(slide)
        return slide

    def new(self, title: str, kind: str = "CONTENT"):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts.get_by_name(kind))
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        title_xml, number_xml = (copy.deepcopy(el) for el in self.proto[kind])
        runs = title_xml.findall(".//" + qn("a:r"))
        runs[0].find(qn("a:t")).text = title
        for extra in runs[1:]:
            extra.getparent().remove(extra)
        slide.shapes._spTree.append(title_xml)
        slide.shapes._spTree.append(number_xml)
        self.order.append(slide)
        return slide

    def save(self, path: str) -> None:
        id_list = self.prs.slides._sldIdLst
        elements = {el.id: el for el in id_list}
        wanted = [s.slide_id for s in self.order]       # looked up before the list is rebuilt
        for el in list(id_list):
            id_list.remove(el)
            if el.id not in wanted:
                self.prs.part.drop_rel(el.rId)
        for slide_id in wanted:
            id_list.append(elements[slide_id])
        self.prs.save(path)


def retext(slide, old: str, new: str) -> None:
    hits = 0
    for sh in slide.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        hits += 1
    if not hits:
        raise ValueError(f"text to replace not found: {old!r}")


def load_icons(prs) -> dict[str, bytes]:
    blobs = {}
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.shape_type is not None and "PICTURE" in str(sh.shape_type) and sh.width < 914400:
                blobs[hashlib.md5(sh.image.blob).hexdigest()[:8]] = sh.image.blob
    return blobs


# ---------------------------------------------------------------------- #
# results
# ---------------------------------------------------------------------- #
RUN_DIRS = {"rule": "output_rule_only", "gemini": "output", "groq": "output_groq", "ollama": "output_ollama"}
RUN_LABELS = {"rule": "Rules only", "gemini": "Rules + Gemini", "groq": "Rules + Groq", "ollama": "Rules + Ollama"}
RUN_TIME = {"gemini": "43 min", "groq": "42 min", "ollama": "2 h 14 min"}   # wall-clock time of the full run
RUN_WHERE = {"gemini": "Google cloud", "groq": "Groq cloud", "ollama": "This laptop (CPU)"}
RUN_COLOURS = {"rule": HEX["a5"], "gemini": HEX["a1"], "groq": HEX["a4"], "ollama": HEX["a2"]}
CATEGORIES = [("arithmetic", "Arithmetic"), ("nested", "Nested"), ("repeated_expr", "Repeated expr."),
              ("dead_code", "Dead code"), ("loops", "Loops"), ("conditionals", "Conditionals"), ("mixed", "Mixed")]
PENDING = "run in progress"


def load_run(key: str, baseline=None) -> dict | None:
    folder = os.path.join(ROOT, RUN_DIRS[key])
    try:
        records = load_results(folder)
    except FileNotFoundError:
        return None
    run = {"records": records, "m": aggregate_metrics(records), "calls": 0, "folder": folder}
    for log in glob.glob(os.path.join(folder, "logs", "*.log")):
        with open(log, encoding="utf-8") as fh:
            run["calls"] += sum(1 for line in fh if line.lstrip().startswith("[LLM]"))
    run["models"] = sorted({r.get("llm_model", "") for r in records if r.get("llm_model")})
    if baseline is not None:
        run["cmp"] = compare_runs(baseline["records"], records)
    return run


def pct(value, digits=2) -> str:
    return "n/a" if value is None else f"{value:.{digits}f}%"


def num(value) -> str:
    return f"{value:,}"


# ---------------------------------------------------------------------- #
# slide content
# ---------------------------------------------------------------------- #
CODE_ORCHESTRATOR = '''
# src/agents/orchestrator.py  (trimmed)
for iteration in range(1, self.max_iterations + 1):
    # 1. ANALYZE (leaving out opportunities already turned down)
    opportunities = [o for o in self.analyzer.analyze(current)
                     if self._opportunity_key(o) not in attempted]
    logger.log_analysis(iteration, opportunities)

    # 2-5. PROPOSE -> VERIFY -> EVALUATE -> ACCEPT/REJECT
    tried, outcome = [], None
    while outcome not in ("accepted", "unverifiable"):
        decision = specialist.route(current, opportunities, llm_calls, tried)
        if decision.route is None:
            break                       # nobody can propose anything more
        tried.append(decision.route)
        proposal = specialist.propose(current, decision, opportunities)
        outcome = self._judge(iteration, current, proposal, logger, proposals)
        if outcome == "accepted":
            current = proposal.candidate
            attempted.clear()           # new opportunities may appear

    if outcome != "accepted" and not opportunities:
        break                           # converged: nothing left to try
'''

CODE_JUDGE = '''
# src/agents/orchestrator.py  (trimmed)
def _judge(self, iteration, current, proposal, logger, proposals):
    # 3. VERIFY: the same gate for rule-based and LLM proposals
    ver_result = self.verifier.verify(current, proposal.candidate)
    logger.log_verification(iteration, ver_result)
    if not ver_result.passed:
        logger.log_reject(iteration, "Verification failed")
        return "invalid"

    # 4. EVALUATE
    old_cost = self.cost_evaluator.evaluate(current)
    new_cost = self.cost_evaluator.evaluate(proposal.candidate)
    logger.log_cost(iteration, old_cost, new_cost)

    # 5. ACCEPT / REJECT: the cost must strictly drop
    if not self.cost_evaluator.is_improvement(old_cost, new_cost):
        logger.log_reject(iteration, "No cost improvement")
        return "rejected"
    logger.log_accept(iteration, "Cost reduced")
    return "accepted"
'''

CODE_COST = '''
# src/agents/cost_evaluator.py
WEIGHTS = {
    "instruction_count": 1.0,
    "arithmetic_ops":    2.0,
    "temp_vars":         0.5,
    "memory_ops":        1.5,
    "branch_ops":        1.0,
}

def is_improvement(self, old, new):
    return (new.weighted_total
            < old.weighted_total)
'''

CODE_ROUTE = '''
# src/agents/optimization_specialist.py  (trimmed)
def route(self, program, opportunities, llm_calls=0, tried=()):
    """Decide who proposes next for ``program``."""
    order = ([Route.RULE, Route.LLM] if self.mode != "llm"
             else [Route.LLM, Route.RULE])
    if not self.llm_attached:
        order = [Route.RULE]
    passed_over, llm_note = [], ""
    for route in order:
        if route in tried:
            name = "rule-based" if route is Route.RULE else "LLM"
            passed_over.append(f"the {name} proposal was turned down")
        elif route is Route.RULE and not opportunities:
            passed_over.append("the rules have nothing left to try")
        elif route is Route.LLM and (
                llm_note := self._llm_blocker(program, llm_calls)):
            passed_over.append(f"LLM not used, {llm_note}")
        else:
            reason = self._reason(route, opportunities, passed_over)
            return RoutingDecision(route, reason, llm_note)
    return RoutingDecision(None, "; ".join(passed_over), llm_note)
'''

CODE_ANALYZER = '''
# src/agents/code_analyzer.py  (trimmed)
def _strength_reduction(self, instrs):
    out = []
    for i, ins in enumerate(instrs):
        if ins.op != TACOp.BINARY or ins.operator != "*":
            continue
        if is_literal(ins.arg2) and is_variable(ins.arg1):
            var, const = ins.arg1, int(ins.arg2)
        elif is_literal(ins.arg1) and is_variable(ins.arg2):
            var, const = ins.arg2, int(ins.arg1)
        else:
            continue
        if const < 2 or const & (const - 1):
            continue                  # not a power of two
        k = const.bit_length() - 1
        if k > MAX_SHIFT:
            continue                  # x << k would be an invalid shift
        if k == 1:
            new = {"operator": "+", "arg1": var, "arg2": var}   # x*2 -> x+x
        else:
            new = {"operator": "<<", "arg1": var, "arg2": str(k)}
        out.append(_make(OpportunityType.STRENGTH_REDUCTION, i, ...))
    return out
'''

CODE_PROMPT = '''
You are a compiler optimization expert. Given this Three-Address Code (TAC),
propose exactly ONE semantics-preserving optimization.
Original TAC:
{tac_string}
Rules:
1. The optimized code MUST produce identical output for ALL possible inputs.
2. Apply exactly one optimization: constant folding, dead code elimination,
   common subexpression elimination, algebraic simplification, strength
   reduction, or copy propagation.
3. Return ONLY the modified TAC inside <tac> tags, nothing else outside.
4. Explain your optimization in one line inside <reasoning> tags.
TAC syntax (one instruction per line, keep exactly this syntax):
  x = y            x = a op b   (op: + - * / % << == != < > <= >= && ||)
  x = -a           x = !a       L0:      goto L0
  if a goto L0     ifFalse a goto L0     print a
Variables that are read before being written are program inputs. Integer
division truncates toward zero; dividing by zero is a runtime error.
Format:
<reasoning>your explanation</reasoning>
<tac>
modified TAC here
</tac>
'''

CODE_LLM = '''
# src/agents/optimizer_llm.py  (trimmed)
def propose(self, program):
    if not self.available:
        return None, f"LLM unavailable: {self.unavailable_reason}"
    try:
        text = self._complete(self._build_prompt(program))
    except _LLMUnavailable as exc:      # quota spent, server gone
        self._disable(str(exc))
        return None, f"LLM unavailable: {self.unavailable_reason}"
    except _LLMSkip as exc:             # only this call failed
        return None, str(exc)
    return self._parse_text(text, program)

def _parse_text(text, original):
    tac_match = _TAC_RE.search(text)
    if not tac_match:
        return None, "no <tac> block in LLM response"
    candidate = parse_llm_tac(tac_match.group(1))
    if candidate is None:
        return None, "could not parse LLM TAC"
    if candidate == original:
        return None, "LLM returned unchanged TAC"
    return candidate, reasoning
'''

CODE_VERIFIER = '''
# src/agents/verifier.py  (trimmed)
def input_sets(self, variables, constants=None):
    rng = random.Random(self.seed)
    lo, hi = self.value_range                      # -100 .. 100
    sets = [{v: rng.randint(lo, hi) for v in variables}
            for _ in range(self.num_random)]       # 20 random sets
    for edge in (0, 1, -1, 99, -99):               # 5 edge cases
        sets.append({v: edge for v in variables})
    pool = []
    if variables and constants:                    # c-1, c, c+1
        pool = sorted({k for c in sorted(constants)
                       for k in (c - 1, c, c + 1)})
        sets.extend({v: k for v in variables} for k in pool)
    sets.extend(self._small_scope_sets(variables, rng))
    sets.extend(self._pairwise_sets(variables, pool))
    return sets

def _small_scope_sets(self, variables, rng):
    for values in SMALL_SCOPES:                    # widest range first
        if len(values) ** len(variables) <= self.max_small_scope:
            return [dict(zip(variables, combo)) for combo in
                    itertools.product(values, repeat=len(variables))]
'''

DEMO_C = '''
#include <stdio.h>

/* A loop with work that does not need
   to be redone every time. */
int main() {
    int sum = 0;
    for (int i = 0; i < 10; i++) {
        int step = 2 + 3;
        sum += step * 1;
    }
    printf("%d\\n", sum);
    return 0;
}
'''

DEMO_TAC = [
    "t0 = 0", "sum = t0", "t1 = 0", "i = t1", "L0:", "t2 = 10", "t3 = i < t2", "ifFalse t3 goto L1",
    "t4 = 2", "t5 = 3", "t6 = t4 + t5", "step = t6", "t7 = 1", "t8 = step * t7", "t9 = sum + t8",
    "sum = t9", "t10 = 1", "t11 = i + t10", "i = t11", "goto L0", "L1:", "print sum"]

DEMO_LOOP = '''
STAGE 3  Agent loop: Analyze -> Propose -> Verify -> Evaluate -> Accept/Reject
  Routing: hybrid (rules first, the LLM when they are turned down or run out)
  LLM: groq openai/gpt-oss-120b
  [ 1] COPY_PROPAGATION   t0 = 0 -> uses of t0 -> 0 at instr 1
         verified  cost 56.0 -> 53.0   ACCEPTED
  [ 6] CONSTANT_FOLD      t6 = 2 + 3 -> t6 = 5
         verified  cost 41.0 -> 40.5   ACCEPTED
  ...  steps 2 to 13: copy propagation and folding, each verified and accepted
  [14] COPY_PROPAGATION   t11 = i + 1; i = t11 -> i = i + 1
         verified  cost 22.5 -> 19.5   ACCEPTED
       No more opportunities found.
       route to LLM optimizer (groq openai/gpt-oss-120b): the rules have nothing left to try
  [15] LLM (Groq)
         Algebraic simplification: the loop adds a constant 5 exactly ten times,
         so replace it with sum = 5*10 = 50 and eliminate the dead loop.
         verified  cost 19.5 -> 3.5    ACCEPTED
  [16] COPY_PROPAGATION   sum = 50 -> uses of sum -> 50 at instr 1
         verified  cost 3.5 -> 1.0     ACCEPTED
       LLM not consulted: program is only constant prints (provably minimal)
  finished in 17 iterations (0.9 s)
'''

DEMO_PROOF = '''
STAGE 6  Proof: the optimized program behaves the same
  original : prints [50]
  optimized: prints [50]
  identical output
  verifier, final check against the original: PASS
'''

LOG_EXAMPLE = '''
=== Program: loops_001.src ===
--- Iteration 20 ---
  [ANALYZE] Found 1 opportunity: COPY_PROPAGATION (priority=50)
  [ROUTE]   Rule-based optimizer: highest-priority opportunity is COPY_PROPAGATION
  [PROPOSE] COPY_PROPAGATION at instr 10: t15 = acc -> uses of t15 -> acc
  [VERIFY]  PASS (differential: 25/25 inputs matched)
  [COST]    Before: 23.0, After: 20.0 (reduction: 13.0%)
  [ACCEPT]  Change accepted
--- Iteration 21 ---
  [ANALYZE] No more opportunities found.
  [ROUTE]   LLM optimizer (groq openai/gpt-oss-120b): the rules have nothing left to try
  [LLM]     Dead code elimination: the loop only assigns the constant 42 to acc
            repeatedly, so it can be replaced by a single assignment.
  [VERIFY]  PASS (differential: 25/25 inputs matched)
  [COST]    Before: 20.0, After: 4.5 (reduction: 77.5%)
  [ACCEPT]  Change accepted
--- Iteration 23 ---
  [ANALYZE] No more opportunities found.
  [INFO]    LLM not consulted: program is only constant prints (provably minimal)
=== RESULT === 72.5 -> 2.0 (97.2% reduction)   Output match: PASS
'''

LOG_ORIGINAL = ["t0 = 0", "acc = t0", "t1 = 0", "i = t1", "L0:", "t2 = 11", "t3 = i < t2", "ifFalse t3 goto L1",
                "t4 = 2", "t5 = 1", "t6 = t4 + t5", "h = t6", "...  (16 more)"]

ROUTE_LOG = '''
--- Iteration 21 ---
  [ANALYZE] No more opportunities found.
  [ROUTE]   LLM optimizer (groq openai/gpt-oss-120b): the rules have nothing left to try
  [LLM]     LLM proposal: Dead code elimination: the loop only assigns the constant 42 to acc repeatedly ...
  [VERIFY]  PASS (differential: 25/25 inputs matched, symbolic: skipped (control flow))
  [COST]    Before: 20.0, After: 4.5 (reduction: 77.5%)
  [ACCEPT]  Change accepted
'''


def notes_panel(slide, x, y, w, h, heading, bullets, size=14, after=9):
    """A light card with a heading and spaced bullet points."""
    card(slide, x, y, w, h)
    text(slide, x + 0.25, y + 0.22, w - 0.5, 0.4, heading, size=18, color="a1", bold=True)
    text(slide, x + 0.25, y + 0.75, w - 0.5, h - 0.95, bullets, size=size, color="tx1", after=after, bullets=True)


def check_row(slide, x, y, heading, detail, w=5.1, icon_name="check", fill="a2"):
    icon(slide, icon_name, x, y + 0.08, d=0.6, fill=fill)
    text(slide, x + 0.82, y, w, 0.42, heading, size=17, color="tx1", bold=True)
    text(slide, x + 0.82, y + 0.44, w, 0.7, detail, size=14)


def stat_card(slide, x, y, w, h, big, label, colour="a1", big_size=40):
    card(slide, x, y, w, h, fill=colour, alpha=10)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.85, big, size=big_size, color=colour, bold=True, anchor="ctr")
    text(slide, x + 0.25, y + 1.1, w - 0.5, h - 1.2, label, size=14)


def dark_card(slide, x, y, w, h, big, label, colour="a6"):
    card(slide, x, y, w, h, fill="bg1", alpha=9)
    text(slide, x + 0.25, y + 0.15, w - 0.5, 0.75, big, size=34, color=colour, bold=True)
    text(slide, x + 0.25, y + 0.95, w - 0.5, h - 1.05, label, size=15, color="bg1")


def column_chart(slide, x, y, w, h, categories, series, labels=False):
    """A native clustered column chart. series: [(name, values, hex colour)], values in percent."""
    data = CategoryChartData()
    data.categories = categories
    for name, values, _ in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(x), E(y), E(w), E(h), data).chart
    chart.font.size = Pt(12)
    chart.font.name = "Calibri"
    chart.font.color.rgb = RGBColor.from_string(HEX["tx2"])
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.overlap = -8
    for plotted, (_, _, colour) in zip(plot.series, series):
        plotted.format.fill.solid()
        plotted.format.fill.fore_color.rgb = RGBColor.from_string(colour)
    if labels:
        plot.has_data_labels = True
        plot.data_labels.number_format = "0.0"
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(10)
    value_axis = chart.value_axis
    value_axis.minimum_scale, value_axis.maximum_scale, value_axis.major_unit = 0, 100, 20
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("E2E8F0")
    value_axis.format.line.fill.background()
    value_axis.tick_labels.number_format = '0"%"'
    value_axis.tick_labels.number_format_is_linked = False
    chart.category_axis.format.line.color.rgb = RGBColor.from_string(BORDER)
    chart.category_axis.has_major_gridlines = False
    return chart


def what_it_shows(slide, bullets):
    card(slide, 9.75, TOP, 2.98, 5.3)
    text(slide, 10.0, 1.65, 2.48, 0.4, "What it shows", size=18, color="a1", bold=True)
    text(slide, 10.0, 2.2, 2.53, 4.3, bullets, size=16, color="tx1", after=14, bullets=True)


def chart_slide(deck, title, image_path, bullets):
    """Same frame as the Review 1 chart slides: picture on the left, 'What it shows' on the right."""
    slide = deck.new(title)
    if image_path and os.path.exists(image_path):
        picture(slide, image_path, LM, TOP, 8.85, 5.3)
    else:
        card(slide, LM, TOP, 8.85, 5.3, fill="bg2")
        text(slide, LM, TOP, 8.85, 5.3, "Chart is added when the run finishes", size=18, color="a5",
             align="ctr", anchor="ctr")
    what_it_shows(slide, bullets)
    return slide


def diagram(name: str) -> str:
    """The slide-ready diagram from this folder if it has been rendered, else the one in the project root."""
    local = os.path.join(HERE, f"review2_{name}.png")
    fallback = {"class_diagram": "class_diagram.png", "data_flow": "data_flow_diagram.png"}[name]
    return local if os.path.exists(local) else os.path.join(ROOT, fallback)


AGENTS = [
    ("1  Code Analysis Specialist", "code_analyzer.py", "the current TAC program",
     ["Builds the control-flow graph",
      "Liveness and available-copies dataflow analyses",
      "Six detectors: constant fold, algebraic, dead code, CSE, strength reduction, copy propagation",
      "Leaves out opportunities already turned down"],
     "opportunities ranked by priority, 100 down to 50"),
    ("2  Optimization Specialist", "optimization_specialist.py", "the program and its opportunities",
     ["Routes the step by mode: rule, hybrid or llm",
      "Option A: the rule-based optimizer applies the top opportunity to a copy",
      "Option B: an LLM (Gemini, Groq or Ollama) answers one shared prompt",
      "The reply is parsed back into TAC; unusable replies are dropped"],
     "one candidate program with its reasoning"),
    ("3  Verification Module", "verifier.py", "current program + candidate",
     ["Runs both on random, edge, boundary, small-scope and pairwise inputs",
      "Compares printed output and failure mode",
      "Symbolic check for straight-line code",
      "Both timing out: cannot be verified, so rejected"],
     "PASS or FAIL, with the reason"),
    ("4  Cost Evaluator", "cost_evaluator.py", "program + verified candidate",
     ["Counts instructions, arithmetic ops, temporaries, copies and branches",
      "Weights them 1.0, 2.0, 0.5, 1.5 and 1.0",
      "Compares the two weighted totals",
      "A change counts only if the cost strictly drops"],
     "ACCEPT, or REJECT for no cost gain"),
]


def architecture_slide(deck):
    """The whole system on one slide, with what happens inside every agent."""
    s = deck.new("System Architecture")
    # front end: source to TAC
    stages = [("Source program", False), ("Lexer", True), ("Parser", True), ("AST", True),
              ("TAC Generator", True), ("TAC program", False)]
    bw, gap, y0, bh = 1.85, 0.206, 1.3, 0.46
    for i, (label, process) in enumerate(stages):
        x = LM + i * (bw + gap)
        box(s, x, y0, bw, bh, label, fill="bg2" if process else "bg1", line="a1" if process else "a5", bold=process)
        if i:
            hline(s, x - gap, y0 + bh / 2, x, y0 + bh / 2, color="a5", width=1.25, arrow=True)
    # the TAC program enters the loop at the first agent
    top, left = 2.06, LM
    hline(s, 11.8, y0 + bh, 11.8, 1.91, color="a5", width=1.25)
    hline(s, 11.8, 1.91, 2.2, 1.91, color="a5", width=1.25)
    hline(s, 2.2, 1.91, 2.2, 2.72, color="a5", width=1.25, arrow=True)
    # the orchestrator frames the four agents
    frame = card(s, left, top, CW, 4.68, fill="a1", alpha=5, line="a1", line_w=1.25, radius=0.1)
    dash = OxmlElement("a:prstDash")
    dash.set("val", "dash")
    frame.line._get_or_add_ln().append(dash)
    text(s, left + 2.0, top + 0.1, CW - 2.25, 0.5, [[
        ("Orchestrator ", {"bold": True, "color": "a1"}), ("(orchestrator.py)  ", {"font": "Consolas", "size": 10.5, "color": "a1"}),
        "runs the loop for up to 50 iterations, remembers what was already tried, logs every decision "
        "and re-verifies the final program against the original"]], size=12, color="tx1")
    aw, agap, ay, ah = 2.72, 0.25, 2.72, 2.98
    for i, (name, source, given, steps, result) in enumerate(AGENTS):
        x = left + 0.25 + i * (aw + agap)
        card(s, x, ay, aw, ah, fill="bg1", line="a1", line_w=1.25, radius=0.07)
        text(s, x + 0.13, ay + 0.09, aw - 0.26, 0.27, name, size=13, color="tx1", bold=True)
        text(s, x + 0.13, ay + 0.36, aw - 0.26, 0.2, source, size=10, color="a1", font="Consolas")
        text(s, x + 0.13, ay + 0.6, aw - 0.26, 0.22, [[("In: ", {"bold": True, "color": "a5"}), given]], size=10.5, color="tx2")
        text(s, x + 0.13, ay + 0.98, aw - 0.26, 1.5, steps, size=10.5, color="tx1", after=2, bullets=True)
        text(s, x + 0.13, ay + ah - 0.42, aw - 0.26, 0.38, [[("Out: ", {"bold": True, "color": "a2"}), result]],
             size=10.5, color="tx1", anchor="b")
        if i:
            hline(s, x - agap, ay + ah / 2, x, ay + ah / 2, color="a5", width=1.5, arrow=True)
    # decision and outputs
    dy = ay + ah + 0.18
    card(s, left + 0.25, dy, CW - 0.5, 0.74, fill="a2", alpha=10)
    text(s, left + 0.45, dy, CW - 0.9, 0.74, [
        [("5  Decision.  ", {"bold": True, "color": "a2"}), ("ACCEPT: ", {"bold": True}),
         "the candidate becomes the current program and the loop analyzes it again.  ",
         ("REJECT: ", {"bold": True}), "the next opportunity is tried, or the same step goes to the other proposer.  ",
         ("When nothing is left: ", {"bold": True}),
         "final check against the original, then the outputs: optimized TAC, decision log, metrics and 7 charts."]],
         size=11, color="tx1", anchor="ctr")
    # from the cost evaluator down to the decision, and back up to the analyzer
    last_x = left + 0.25 + 3 * (aw + agap) + aw / 2
    hline(s, last_x, ay + ah, last_x, dy, color="a2", width=1.5, arrow=True)
    hline(s, left + 0.25 + aw / 2, dy, left + 0.25 + aw / 2, ay + ah, color="a2", width=1.5, arrow=True)
    return s


def code_slide(deck, title, source, heading, bullets, size=12, lang="py"):
    slide = deck.new(title)
    code(slide, LM, TOP, 8.0, 5.35, source, size=size, lang=lang)
    notes_panel(slide, 8.85, TOP, 3.88, 5.35, heading, bullets)
    return slide


# ---------------------------------------------------------------------- #
def build() -> str:
    global ICON_BLOBS
    base = find_base()
    deck = Deck(base)
    ICON_BLOBS = load_icons(deck.prs)

    rule = load_run("rule")
    runs = {key: load_run(key, rule) for key in ("gemini", "groq", "ollama")}
    gemini, groq, ollama = runs["gemini"], runs["groq"], runs["ollama"]
    total_tests = 336

    def avg(run, digits=1):
        return pct(run["m"]["avg_cost_reduction_pct"], digits) if run else "pending"

    # 1 ---------------------------------------------------------------- title (team line updated)
    s = deck.keep(1)
    team = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Team:"))
    para = team.text_frame.paragraphs[0]._p
    names = para.findall(qn("a:r"))
    names[1].find(qn("a:t")).text = "ADARSH SHARMA BADHAI[24BCE2972]"
    names[2].find(qn("a:t")).text = "  ·  NIRUPAM RAUT[24BAI0345]"
    for element in list(para):           # drop the line break and the underlined repository link after it
        underlined = element.tag == qn("a:r") and element.find(qn("a:rPr")).get("u") == "sng"
        if element.tag == qn("a:br") or underlined:
            para.remove(element)

    # 2 ---------------------------------------------------------------- objectives
    s = deck.new("Objectives")
    objectives = [
        ("Agent architecture", "Five components as separate classes with clear interfaces"),
        ("Orchestrator loop", "Analyze → Propose → Verify → Evaluate → Accept/Reject"),
        ("6 optimization types", "Folding, algebraic, dead code, CSE, strength reduction, copy propagation"),
        ("Dual verification", "Differential testing on hundreds of input sets + symbolic equivalence"),
        ("Cost-gated acceptance", "Accept only if verified correct AND cheaper"),
        ("LLM proposers", "Gemini, Groq and a local Ollama model behind one router"),
        ("Offline mode", "The whole agent runs with no internet connection"),
        ("Full evaluation", "500 programs, the 6 required metrics, 7 charts, 4 configurations"),
    ]
    for i, (heading, detail) in enumerate(objectives):
        check_row(s, LM + (i // 4) * 6.2, 1.45 + (i % 4) * 1.32, heading, detail)

    # 4-7 -------------------------------------------------------------- kept from Review 1
    deck.keep(4)
    deck.keep(5)
    s = deck.keep(6)
    retext(s, "ask Gemini if rules fail or run out", "or ask the LLM, as the router decides")
    deck.keep(7)

    # 8 ---------------------------------------------------------------- architecture (redrawn in detail)
    architecture_slide(deck)

    # 9 ---------------------------------------------------------------- data flow diagram
    chart_slide(deck, "Data Flow Diagram", diagram("data_flow"), [
        "Source code becomes a TAC program in the front end",
        "The Code Analyzer sends ranked opportunities to the Optimization Specialist",
        "The Specialist routes the step: Option A the rules, Option B an LLM",
        "Each candidate passes the Verifier, then the Cost Evaluator"])

    # 10 --------------------------------------------------------------- class diagram
    chart_slide(deck, "Class Diagram", diagram("class_diagram"), [
        "The Orchestrator owns the four other components and writes to the Decision Logger",
        "The Optimization Specialist owns the rule-based optimizer and holds one LLM backend",
        "Gemini, Groq and Ollama share one base class: same prompt, same reply parser"])

    # 11 --------------------------------------------------------------- workflow (kept, updated)
    s = deck.keep(9)
    retext(s, "(rules or Gemini)", "(rules or an LLM)")
    retext(s, "Gemini is asked when", "In hybrid mode the LLM is asked when")

    # 12 --------------------------------------------------------------- routing
    s = deck.new("Routing: Rule-Based or LLM")
    modes = [("wrench", "rule", "Rule-based only", "Always the highest-priority opportunity from the analyzer. The default: instant and deterministic."),
             ("route", "hybrid", "Rules first, LLM second", "The LLM is asked when a rule proposal is turned down, or when the rules have nothing left to try."),
             ("chat", "llm", "LLM first, rules second", "The rules take over when the LLM is not running, out of call budget, or its proposal is turned down.")]
    for i, (name, mode, heading, detail) in enumerate(modes):
        x = LM + i * 4.127
        card(s, x, TOP, 3.88, 2.5)
        icon(s, name, x + 0.22, TOP + 0.22)
        text(s, x + 0.95, TOP + 0.2, 2.7, 0.35, [[("--optimizer ", {"color": "a5"}), (mode, {"color": "a1", "bold": True})]],
             size=14, font="Consolas")
        text(s, x + 0.95, TOP + 0.52, 2.7, 0.4, heading, size=17, color="tx1", bold=True)
        text(s, x + 0.22, TOP + 1.08, 3.44, 1.35, detail, size=14)
    text(s, LM, 4.08, CW, 0.35, "Every routing decision is written to the decision log with its reason",
         size=15, color="tx1", bold=True)
    code(s, LM, 4.5, CW, 2.25, ROUTE_LOG, size=12.5, lang="log", pad=0.24)

    # 13 --------------------------------------------------------------- LLM backends
    s = deck.new("Three LLM Backends, One Interface")
    backends = [("Gemini", "gemini-3.1-flash-lite", ["Runs in Google's cloud", "Free API key (.gemini_key)", "Called through the google-genai SDK",
                                                    "Strongest results so far: " + avg(gemini)]),
                ("Groq", "openai/gpt-oss-120b", ["Runs in Groq's cloud", "Free API key (.groq_key)", "Plain HTTP, no SDK to install",
                                                 "Moves to the next model when a daily quota is spent"]),
                ("Ollama", "qwen2.5-coder:7b", ["Runs on this laptop", "No key and no internet", "Plain HTTP to the local server",
                                                "Falls back to the CPU if the GPU runner crashes"])]
    for i, (name, model, points) in enumerate(backends):
        x = LM + i * 4.127
        card(s, x, TOP, 3.88, 3.45)
        icon(s, "chat" if i < 2 else "layers", x + 0.22, TOP + 0.22)
        text(s, x + 0.95, TOP + 0.18, 2.7, 0.42, name, size=20, color="tx1", bold=True)
        text(s, x + 0.95, TOP + 0.6, 2.75, 0.3, model, size=12.5, color="a1", font="Consolas", bold=True)
        text(s, x + 0.25, TOP + 1.15, 3.4, 2.2, points, size=15, color="tx1", after=9, bullets=True)
    card(s, LM, 5.15, CW, 1.6, fill="a1", alpha=10)
    icon(s, "route", LM + 0.25, 5.67)
    text(s, LM + 1.05, 5.3, CW - 1.3, 1.3, [
        [("--llm-provider auto ", {"font": "Consolas", "bold": True, "color": "a1"}),
         "tries Gemini, then Groq, then Ollama, and uses the first one that is configured and reachable. ",
         ("--offline ", {"font": "Consolas", "bold": True, "color": "a1"}), "always uses Ollama."],
        "All three share one prompt, one reply parser and the same verification and cost gates."],
        size=14, color="tx1", after=6, anchor="ctr")

    # 14 --------------------------------------------------------------- tools
    s = deck.new("Implementation: Tools & Technologies")
    table(s, LM, TOP, [1.55, 2.75, 3.3], [
        ["Layer", "Technology", "Purpose"],
        ["Language", "Python 3.10+", "match statements, dataclasses, type hints"],
        ["Cloud LLMs", "Gemini, Groq", "Free-tier proposers (SDK and plain HTTP)"],
        ["Local LLM", "Ollama + qwen2.5-coder:7b", "Offline proposer, no key needed"],
        ["Visualization", "Matplotlib, Seaborn, Plotly", "Bar, line, scatter, heatmap, pie, Sankey"],
        ["Testing", f"pytest ({total_tests} tests)", "Unit, integration, stand-in LLM servers"],
        ["Live demo", "demo.py (C subset)", "Walks a real C program through the loop"],
        ["Numerical", "NumPy", "Trend lines and statistics"],
    ], size=14, row_h=0.66)
    card(s, 8.45, TOP, 4.28, 5.28)
    text(s, 8.7, 1.62, 3.8, 0.4, "Project structure", size=18, color="tx1", bold=True)
    structure = [("src/frontend/", "Lexer, Parser, AST, TAC generator, C front end"),
                 ("src/tac/", "Instructions, interpreter, dataflow analysis"),
                 ("src/agents/", "Orchestrator, Code Analyzer, Optimization Specialist, rule + LLM optimizers, Verifier, Cost Evaluator, Decision Logger"),
                 ("src/visualization/", "7 charts, metrics report"),
                 ("tests/", f"{total_tests} tests, one module per component")]
    text(s, 8.7, 2.15, 3.8, 4.4, [[(folder + "  ", {"font": "Consolas", "bold": True, "color": "a1", "size": 13}), detail]
                                  for folder, detail in structure], size=14, color="tx2", after=9)

    # 15 --------------------------------------------------------------- the five components
    s = deck.new("Implementation: The Five Components")
    components = [
        ("redo", "1. Orchestrator", "Runs the loop, logs every decision, stops when nothing is left (or after 50 iterations), then re-verifies the result against the original"),
        ("search", "2. Code Analysis Specialist", "Dataflow analysis (CFG, liveness, available copies). Ranks 6 types: Fold 100 → Algebraic 90 → Dead Code 80 → CSE 70 → Strength 60 → Copy Prop 50"),
        ("route", "3. Optimization Specialist", "Routes each step to the rule-based optimizer or to an LLM (Gemini, Groq or Ollama) and returns one candidate program"),
        ("shield", "4. Verification Module", "Differential testing on random, edge, boundary, small-scope and pairwise inputs, plus a symbolic check for straight-line code"),
        ("scale", "5. Cost Evaluator", "cost = 1.0·instructions + 2.0·arithmetic + 0.5·temps + 1.5·copies + 1.0·branches. A change must strictly lower it"),
    ]
    spots = [(LM, TOP), (LM + 4.127, TOP), (LM + 8.254, TOP), (2.66, 4.13), (6.79, 4.13)]
    for (name, heading, detail), (x, y) in zip(components, spots):
        card(s, x, y, 3.88, 2.5)
        icon(s, name, x + 0.22, y + 0.22)
        text(s, x + 0.95, y + 0.3, 2.85, 0.45, heading, size=17, color="tx1", bold=True)
        text(s, x + 0.22, y + 0.95, 3.46, 1.5, detail, size=14)

    # 16-22 ------------------------------------------------------------ code
    code_slide(deck, "Code: The Orchestrator Loop", CODE_ORCHESTRATOR, "How it works", [
        "One iteration is one step: analyze, propose, verify, evaluate, decide",
        "The specialist picks the route; if that proposal is turned down it may hand the same step to the other proposer",
        "An accepted change becomes the current program and clears the \"already tried\" list",
        "The loop ends when nothing was accepted and no opportunity is left"])

    s = deck.new("Code: The Verify and Cost Gates")
    code(s, LM, TOP, 7.75, 5.35, CODE_JUDGE)
    code(s, 8.6, TOP, 4.13, 3.1, CODE_COST, size=11.5)
    card(s, 8.6, 4.7, 4.13, 2.05)
    text(s, 8.85, 4.88, 3.63, 1.75, [
        "Rule-based and LLM proposals share this path: an LLM answer gets no shortcut",
        "A correct change that is not cheaper is still rejected"], size=14, color="tx1", after=9, bullets=True)

    code_slide(deck, "Code: The Routing Decision", CODE_ROUTE, "How it works", [
        "The mode sets the order: rules first (hybrid) or LLM first (llm)",
        "A route is skipped if its proposal was just turned down, the rules have nothing, or the LLM is blocked",
        "LLM blockers: not running, call budget used up, or the program is already provably minimal",
        "Each decision carries its reason, which becomes a [ROUTE] line in the log"])

    code_slide(deck, "Code: Finding an Opportunity", CODE_ANALYZER, "How it works", [
        "Each of the six detectors returns Opportunity objects: type, priority, location, description",
        "Here: x × 2 becomes x + x, and x × 2^k becomes a shift",
        "The k > 64 guard came out of fuzz testing: the interpreter rejects larger shifts",
        "The rule-based optimizer then applies exactly one opportunity to a copy of the program"])

    code_slide(deck, "Code: The Prompt Sent to the LLM", CODE_PROMPT, "How it works", [
        "One prompt for all three LLMs, so their results are comparable",
        "It asks for exactly one change, in strict TAC syntax",
        "The reply must carry <reasoning> and <tac> blocks; anything else is discarded",
        "The model's answer is only a suggestion: it still has to pass the verifier",
        "Blank lines are removed here to fit the slide"], size=11.5, lang="txt")

    code_slide(deck, "Code: Calling the LLM", CODE_LLM, "How it works", [
        "A backend only implements _complete(prompt): Gemini through its SDK, Groq and Ollama through plain HTTP",
        "Two kinds of failure: skip this one proposal, or switch the LLM off for the rest of the run",
        "Unparseable or unchanged TAC never reaches the verifier",
        "If the LLM switches off, the run continues rule-based"])

    code_slide(deck, "Code: The Verifier's Input Sets", CODE_VERIFIER, "How it works", [
        "Original and candidate run on every set; printed output and failure mode must match",
        "Small scope: every combination of small values, so one input can be zero while another is not",
        "Pairwise: every two inputs meet every two boundary values",
        "No inputs at all: one execution, so a match is a proof"])

    # 23 --------------------------------------------------------------- verification deep dive
    s = deck.new("Verification Deep Dive")
    table(s, LM, TOP, [1.75, 3.55, 3.2], [
        ["Input sets", "What is tried", "What it catches"],
        ["Random", "20 sets, values in -100..100", "General differences"],
        ["Edge cases", "All inputs 0, 1, -1, 99, -99", "Zero and sign handling"],
        ["Boundary", "c-1, c, c+1 for every literal c", "Off-by-one comparisons"],
        [("Small scope", {"bold": True, "color": "a1"}), "Every combination of small values (one input: -128..128)", "&& against ||, a branch on a flag"],
        [("Pairwise", {"bold": True, "color": "a1"}), "Every two inputs × every two boundary values", "Two inputs that must line up"],
    ], size=14, row_h=0.62)
    text(s, LM, 5.25, 8.5, 0.35, [[("Small scope", {"bold": True, "color": "a1"}), " and ",
                                   ("pairwise", {"bold": True, "color": "a1"}),
                                   " were added after mutation testing exposed a gap."]], size=14)
    gates = [("1", "Differential testing", "Output and failure mode must match on every input set"),
             ("2", "Symbolic check", "Straight-line code: canonical polynomials, so x·2 = x+x = x<<1"),
             ("3", "Cost", "A correct change is still rejected unless the cost strictly drops")]
    for i, (n, heading, detail) in enumerate(gates):
        y = TOP + i * 1.3
        card(s, 9.35, y, 3.38, 1.15)
        card(s, 9.55, y + 0.3, 0.55, 0.55, fill="a1", shape=MSO_SHAPE.OVAL)
        text(s, 9.55, y + 0.3, 0.55, 0.55, n, size=18, color="bg1", bold=True, align="ctr", anchor="ctr")
        text(s, 10.25, y + 0.12, 2.35, 0.35, heading, size=15, color="tx1", bold=True)
        text(s, 10.25, y + 0.45, 2.38, 0.7, detail, size=12)
    card(s, LM, 5.7, CW, 1.05, fill="a2", alpha=12)
    text(s, LM + 0.3, 5.7, CW - 0.6, 1.05, [[("Strength of the guarantee: ", {"bold": True, "color": "a2"}),
         "proven for programs with no inputs and for straight-line code; tested on hundreds of input sets "
         "for branching code with inputs. The finished program is re-verified against the original."]],
         size=14, color="tx1", anchor="ctr")

    # 24 --------------------------------------------------------------- validation
    s = deck.new("Validation: Testing the System Itself")
    checks = [("5,100", "random TAC programs: 20,968 rule proposals checked on inputs the verifier does not use, none unsound", "a1"),
              ("800", "deliberately broken programs: none slips past the verifier now (13 of the first 400 did before)", "a1"),
              ("64,620", "re-executions of the Gemini results on small-scope and random inputs, no mismatch", "a1"),
              ("383", "recorded Gemini replies replayed through the stronger verifier: the same decision every time", "a2")]
    for i, (big, label, colour) in enumerate(checks):
        stat_card(s, LM + i * 3.093, TOP, 2.85, 2.5, big, label, colour)
    text(s, LM, 4.15, CW, 0.4, "Bugs found and fixed", size=18, color="tx1", bold=True)
    fixes = [("shield", "Verifier input gap", "Wrong changes that only showed for certain zero / non-zero input mixes could pass"),
             ("wrench", "Strength reduction", "x × 2^k was rewritten to a shift even for k > 64, which fails at run time"),
             ("layers", "Ollama on old GPU drivers", "The GPU runner crashed on load; the call now repeats on the CPU"),
             ("chart", "Metrics table", "The LLM validity rate read \"n/a %\" when an LLM made no proposals")]
    for i, (name, heading, detail) in enumerate(fixes):
        x, y = LM + (i % 2) * 6.2, 4.7 + (i // 2) * 1.02
        icon(s, name, x, y + 0.06, d=0.5)
        text(s, x + 0.7, y, 5.3, 0.35, heading, size=15, color="tx1", bold=True)
        text(s, x + 0.7, y + 0.36, 5.3, 0.6, detail, size=13)

    # 25-27 ------------------------------------------------------------ live demo
    s = deck.new("Demo: A Real C Program")
    text(s, LM, TOP, 5.9, 0.35, "1  The C program (demo_programs/3_loop.c)", size=15, color="tx1", bold=True)
    code(s, LM, 1.85, 5.9, 3.75, DEMO_C, size=13, lang="c")
    text(s, 6.83, TOP, 5.9, 0.35, "2  Front end: C → AST → Three-Address Code", size=15, color="tx1", bold=True)
    card(s, 6.83, 1.85, 5.9, 3.75, fill=CODE_BG)
    half = (len(DEMO_TAC) + 1) // 2
    for col, chunk in enumerate((DEMO_TAC[:half], DEMO_TAC[half:])):
        lines = [[(f"{col * half + n:>2} | ", {"color": CODE_DIM}), (line if line.endswith(":") else "  " + line, {})]
                 for n, line in enumerate(chunk)]
        text(s, 7.08 + col * 2.75, 2.05, 2.7, 2.6, lines, size=12.5, color=CODE_FG, font="Consolas", after=1)
    hline(s, 7.08, 4.68, 12.48, 4.68, color="374151")
    text(s, 7.08, 4.8, 5.4, 0.7, [[("22 instructions", {"bold": True}), " · 4 arithmetic · 12 temporaries"],
                                   ["12 copies · 2 branches · ", ("weighted cost = 56.0", {"bold": True, "color": CODE_AMBER})]],
         size=12, color=CODE_FG, font="Consolas", after=3)
    card(s, LM, 5.8, CW, 0.95, fill="a1", alpha=10)
    icon(s, "code", LM + 0.22, 6.0)
    text(s, LM + 1.0, 5.8, CW - 1.25, 0.95, [
        [("py demo.py demo_programs/3_loop.c --use-llm --llm-provider groq", {"font": "Consolas", "bold": True, "color": "a1", "size": 13})],
        "The demo walks one program through all six stages, pausing between them. What follows is a real run."],
        size=14, color="tx1", after=3, anchor="ctr")

    s = deck.new("Demo: The Agent at Work")
    code(s, LM, TOP, CW, 5.35, DEMO_LOOP, size=13, lang="log", pad=0.25)

    s = deck.new("Demo: Result and Proof")
    stat_card(s, LM, TOP, 3.88, 1.75, "56.0 → 1.0", "weighted cost", "a1", big_size=36)
    stat_card(s, LM, 3.3, 3.88, 1.75, "98.2%", "cost reduction, 22 instructions down to 1", "a2", big_size=36)
    card(s, LM, 5.2, 3.88, 1.55, fill=CODE_BG)
    text(s, LM + 0.25, 5.33, 3.4, 1.3, [[("Optimized TAC", {"color": CODE_DIM, "bold": True})], [("print 50", {"color": CODE_GREEN, "bold": True, "size": 20})]],
         size=12.5, color=CODE_FG, font="Consolas", after=6)
    text(s, 4.73, TOP, 8.0, 0.35, "4  Suggestions made by the agent", size=15, color="tx1", bold=True)
    table(s, 4.73, 1.85, [2.6, 1.35, 1.35, 1.35, 1.35], [
        ["Source", "Suggested", "Verified", "Accepted", "Invalid"],
        ["Rule-based", "15", "15", "15", "0"],
        ["LLM (Groq)", "1", "1", "1", "0"],
        [("Total", {"bold": True}), ("16", {"bold": True}), ("16", {"bold": True}), ("16", {"bold": True, "color": "a2"}), ("0", {"bold": True})],
    ], size=14, row_h=0.5)
    text(s, 4.73, 4.05, 8.0, 0.35, "6  Proof: both versions are run and compared", size=15, color="tx1", bold=True)
    code(s, 4.73, 4.5, 8.0, 2.25, DEMO_PROOF, size=12.5, lang="log")

    # 28 --------------------------------------------------------------- key metrics
    s = deck.new("Results: Key Metrics")
    order = [("rule", rule), ("gemini", gemini), ("groq", groq), ("ollama", ollama)]

    def cell(run, getter, pending=PENDING, **opts):
        if run is None:
            return (pending, {"color": "a5"})
        value = getter(run["m"])
        return (value, opts) if opts else value

    def llm_cell(key, run, getter):
        if key == "rule":
            return "n/a"
        return cell(run, getter)

    best = max((r["m"]["avg_cost_reduction_pct"] for _, r in order if r), default=0)
    metric_rows = [["Metric"] + [RUN_LABELS[k] for k, _ in order]]
    metric_rows.append(["Average cost reduction"] + [
        (pct(r["m"]["avg_cost_reduction_pct"]), {"bold": r["m"]["avg_cost_reduction_pct"] == best,
                                                 "color": "a1" if r["m"]["avg_cost_reduction_pct"] == best else "tx1"})
        if r else (PENDING, {"color": "a5"}) for _, r in order])
    metric_rows.append(["Programs with >30% reduction"] + [cell(r, lambda m: f"{m['programs_with_reduction_over_30pct']} / {m['total_programs']}") for _, r in order])
    metric_rows.append(["Total proposals"] + [cell(r, lambda m: num(m["total_proposals"])) for _, r in order])
    metric_rows.append(["Accepted"] + [cell(r, lambda m: num(m["accepted"])) for _, r in order])
    metric_rows.append(["Verification pass rate"] + [cell(r, lambda m: pct(m["verification_pass_rate_pct"])) for _, r in order])
    metric_rows.append(["Acceptance rate (of verified)"] + [cell(r, lambda m: pct(m["acceptance_rate_pct"])) for _, r in order])
    metric_rows.append(["Steps to convergence (mean)"] + [cell(r, lambda m: f"{m['steps_to_convergence_mean']:.1f}") for _, r in order])
    metric_rows.append(["LLM validity rate"] + [llm_cell(k, r, lambda m: pct(m["llm_validity_rate_pct"])) for k, r in order])
    metric_rows.append(["Wrong LLM proposals blocked"] + [llm_cell(k, r, lambda m: num(m["by_source"]["LLM"]["invalid"])) for k, r in order])
    metric_rows.append(["False positives"] + [cell(r, lambda m: str(m["false_positive_count"]), bold=True, color="a2") for _, r in order])
    table(s, LM, TOP, [3.73, 2.1, 2.1, 2.1, 2.1], metric_rows, size=14, row_h=0.485)

    # 29 --------------------------------------------------------------- by category
    s = deck.new("Results: Cost Reduction by Category")
    cat_rows = [["Category", "Programs"] + [RUN_LABELS[k] for k, _ in order]]
    for key, label in CATEGORIES:
        row = [label, str(rule["m"]["per_category"][key]["programs"])]
        for _, r in order:
            row.append(pct(r["m"]["per_category"][key]["avg_reduction_pct"], 1) if r else (PENDING, {"color": "a5"}))
        cat_rows.append(row)
    cat_rows.append([("All", {"bold": True}), ("500", {"bold": True})] + [
        (pct(r["m"]["avg_cost_reduction_pct"], 1), {"bold": True}) if r else (PENDING, {"color": "a5"}) for _, r in order])
    table(s, LM, TOP, [2.65, 1.4, 2.02, 2.02, 2.02, 2.02], cat_rows, size=14, row_h=0.5)
    improved = []
    for key in ("gemini", "groq", "ollama"):
        r = runs[key]
        improved.append(f"{RUN_LABELS[key].split('+ ')[1]} " + (str(r["cmp"]["overall"]["programs_improved"]) if r else f"({PENDING})"))
    worse = sum(r["cmp"]["overall"]["programs_worse"] for r in runs.values() if r)
    text(s, LM, 6.05, CW, 0.75, [
        "The rules alone already do well on straight-line code. LLMs matter where the rules cannot remove a loop or a branch.",
        [("Programs improved over rules only:  ", {"bold": True}), "  ·  ".join(improved),
         "  ·  none worse in any run" if not worse else f"  ·  {worse} worse"]], size=14, color="tx1", after=4)

    # 30 --------------------------------------------------------------- native chart
    s = deck.new("Results: Four Configurations Compared")
    shown = [(k, r) for k, r in order if r]
    column_chart(s, LM, TOP, CW, 4.75, [label for _, label in CATEGORIES],
                 [(RUN_LABELS[k], [round(r["m"]["per_category"][c]["avg_reduction_pct"], 1) for c, _ in CATEGORIES],
                   RUN_COLOURS[k]) for k, r in shown])
    takeaway = "Average cost reduction per category. Loops, conditionals and mixed programs are where an LLM adds the most."
    if groq is None:
        takeaway += " The Groq series is added when its run finishes."
    text(s, LM, 6.25, CW, 0.5, takeaway, size=14, color="tx1")

    # 31-37 ------------------------------------------------------------ charts kept from Review 1 (Gemini run)
    for number, title in ((15, "Proposal Flow (Gemini Run)"), (16, "Outcomes by Type (Gemini Run)"),
                          (17, "Accepted Optimizations (Gemini Run)"), (18, "Cost Over Iterations (Gemini Run)"),
                          (19, "Steps vs. Improvement (Gemini Run)"), (20, "Where Optimizations Applied (Gemini Run)")):
        deck.keep(number, "Results: " + title)
    deck.keep(21)

    # 38-39 ------------------------------------------------------------ Groq and Ollama against the rules
    def versus(key, run, extra):
        name = RUN_LABELS[key].split("+ ")[1]
        if run is None:
            return chart_slide(deck, f"Results: Rules Only vs. Rules + {name}", None, [
                "The full Groq run is in progress", "This slide is filled in when it finishes"])
        per = run["m"]["per_category"]
        base = rule["m"]["per_category"]
        bullets = [[(f"{label}: ", {"bold": True}), f"{base[c]['avg_reduction_pct']:.1f}% → {per[c]['avg_reduction_pct']:.1f}%"]
                   for c, label in (("loops", "Loops"), ("conditionals", "Conditionals"), ("mixed", "Mixed"))]
        overall = run["cmp"]["overall"]
        bullets.append([(f"{overall['programs_improved']} programs improved", {"bold": True}),
                        ", none got worse" if not overall["programs_worse"] else f", {overall['programs_worse']} got worse"])
        bullets.extend(extra)
        slide = deck.new(f"Results: Rules Only vs. Rules + {name}")
        column_chart(slide, LM, TOP, 8.85, 5.3, [label for _, label in CATEGORIES],
                     [(RUN_LABELS[k], [round(r["m"]["per_category"][c]["avg_reduction_pct"], 1) for c, _ in CATEGORIES],
                       RUN_COLOURS[k]) for k, r in (("rule", rule), (key, run))], labels=True)
        what_it_shows(slide, bullets)
        return slide

    versus("groq", groq, [])
    versus("ollama", ollama, ["A small local model: modest gains, but with no internet at all"])

    # 40 --------------------------------------------------------------- the three LLMs
    s = deck.new("Results: Comparing the Three LLMs")
    llm_order = [("gemini", gemini), ("groq", groq), ("ollama", ollama)]

    def llm_row(label, getter, **opts):
        row = [label]
        for key, run in llm_order:
            if run is None:
                row.append((PENDING, {"color": "a5"}))
            else:
                value = getter(key, run)
                row.append((value, opts) if opts else value)
        return row

    def model_names(key, run):
        names = [m.split("/")[-1] for m in run["models"]] or [{"gemini": "gemini-3.1-flash-lite"}.get(key, "")]
        return names[0] + (f" (+{len(names) - 1} fallback)" if len(names) > 1 else "")

    llm_rows = [["", "Gemini", "Groq", "Ollama"],
                llm_row("Model", model_names),
                llm_row("Runs on", lambda k, r: RUN_WHERE[k]),
                llm_row("Average cost reduction", lambda k, r: pct(r["m"]["avg_cost_reduction_pct"]), bold=True),
                llm_row("Programs improved over rules", lambda k, r: str(r["cmp"]["overall"]["programs_improved"])),
                llm_row("Calls to the model", lambda k, r: num(r["calls"])),
                llm_row("Usable proposals", lambda k, r: num(r["m"]["by_source"]["LLM"]["suggested"])),
                llm_row("Accepted", lambda k, r: num(r["m"]["by_source"]["LLM"]["accepted"])),
                llm_row("Wrong, blocked by the verifier", lambda k, r: num(r["m"]["by_source"]["LLM"]["invalid"])),
                llm_row("Validity rate", lambda k, r: pct(r["m"]["llm_validity_rate_pct"])),
                llm_row("Full-run time", lambda k, r: RUN_TIME[k] or "see log"),
                llm_row("False positives", lambda k, r: str(r["m"]["false_positive_count"]), bold=True, color="a2")]
    table(s, LM, TOP, [3.7, 2.81, 2.81, 2.81], llm_rows, size=14, row_h=0.445, first_col_bold=False)

    # 41 --------------------------------------------------------------- decision log
    s = deck.new("Example: Agent Decision Log with Routing")
    card(s, LM, TOP, CW, 5.35, fill=CODE_BG)
    text(s, LM + 0.3, TOP + 0.25, 8.75, 4.9, [[(t, {"color": c, "bold": b}) for t, c, b in _log_runs(line)]
                                              for line in LOG_EXAMPLE.strip("\n").split("\n")],
         size=12, color=CODE_FG, font="Consolas", after=2)
    hline(s, 9.85, TOP + 0.3, 9.85, TOP + 5.05, color="374151")
    side = [[("Original TAC (28)", {"color": CODE_DIM, "bold": True})]] + [[line] for line in LOG_ORIGINAL]
    side += [[" "], [("Optimized TAC (2)", {"color": CODE_GREEN, "bold": True})], ["print 42"], ["print 42"]]
    text(s, 10.1, TOP + 0.25, 2.45, 4.9, side, size=12, color=CODE_FG, font="Consolas", after=2)

    # 42 --------------------------------------------------------------- status
    s = deck.new("Project Status")
    columns = [("check", "a2", "System built", ["Front end: C and mini-language to TAC", "Orchestrator loop + decision logging",
                                                "Code Analyzer with 6 optimization types", "Optimization Specialist router",
                                                "Three LLM backends: Gemini, Groq, Ollama", "Offline mode with a local model",
                                                "Live C demo front end"]),
               ("shield", "a1", "Tested and measured", ["Differential + symbolic verification", "Small-scope and pairwise inputs",
                                                        f"{total_tests} automated tests, all passing", "5,100 random programs fuzzed",
                                                        "800 mutation tests: none slips through", "500-program dataset, 4 full runs",
                                                        "7 charts and the 6 required metrics"]),
               ("route", "a4", "Possible next steps", ["Execution time in the cost model", "Prompt tuning for small local models",
                                                      "Ollama on the GPU (driver update)", "Larger, real-world-style programs"])]
    for i, (name, colour, heading, items) in enumerate(columns):
        x = LM + i * 4.127
        card(s, x, TOP, 3.88, 5.35)
        icon(s, name, x + 0.25, TOP + 0.22, d=0.5, fill=colour)
        text(s, x + 0.9, TOP + 0.27, 2.8, 0.42, heading, size=18, color="tx1", bold=True)
        text(s, x + 0.25, TOP + 0.95, 3.4, 4.3, items, size=15, color="tx1", after=11, bullets=True)

    # 43 --------------------------------------------------------------- conclusion
    s = deck.new("Conclusion", kind="DARK")
    conclusion = [(avg(rule), "average cost reduction with the rules alone", "a6"),
                  (avg(gemini), "with Gemini proposals, the best configuration so far", "a6"),
                  (avg(groq) if groq else "Groq", "with Groq's gpt-oss models" if groq else "full run in progress: result added when it finishes", "a6"),
                  (avg(ollama), "with a local model, completely offline", "a6"),
                  ("0", "incorrect changes accepted in any run, re-checked on about 64,000 executions per run", CODE_GREEN),
                  ("13 → 0", "wrong changes slipping past the verifier once mutation testing exposed the gap", "a6")]
    for i, (big, label, colour) in enumerate(conclusion):
        dark_card(s, LM + (i % 3) * 4.127, TOP + (i // 3) * 2.2, 3.88, 1.95, big, label, colour)
    text(s, LM, 5.9, CW, 0.75, [[("Key insight: ", {"bold": True, "color": "a6"}),
                                 "an LLM is a useful proposer whether it runs in the cloud or on a laptop, "
                                 "as long as every suggestion is verified before it is trusted."]], size=18, color="bg1")

    # 44 --------------------------------------------------------------- references (kept, two added)
    s = deck.keep(25)
    column = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("[7]"))
    last = column.text_frame.paragraphs[-1]._p
    for entry in ("[13]  Ollama (2026). Ollama: run large language models locally. https://ollama.com",
                  "[14]  Groq (2026). GroqCloud API documentation. https://console.groq.com/docs"):
        clone = copy.deepcopy(last)
        runs = clone.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = entry
        for extra in runs[1:]:
            clone.remove(extra)
        last.addnext(clone)
        last = clone

    try:
        deck.save(OUTPUT)
        return OUTPUT
    except PermissionError:
        # The deck is open in PowerPoint: leave it alone and write next to it instead.
        fallback = OUTPUT.replace(".pptx", "_new.pptx")
        deck.save(fallback)
        return fallback + "  (the main file is open in PowerPoint; close it and rebuild to replace it)"


if __name__ == "__main__":
    print("saved", build())
