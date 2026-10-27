"""Turn a report .tex into a Word file in the department template's format, with what is new marked in yellow.

The LaTeX is converted with Pandoc, and the result is then brought into the template's format:
A4, 1.5" left margin and 1" elsewhere, Times New Roman 12 pt, 1.5 line spacing, justified text,
16 pt centred upper-case chapter titles on a new page, 14 pt upper-case main headings, 10 pt
captions (below figures, above tables), page numbers centred at the bottom.

Everything that is not in the earlier report is highlighted in yellow. The earlier report is
given as plain text (extracted from its PDF) or as a .docx; a sentence, table cell, caption or
heading counts as old when its letters and digits appear there, whatever the spacing,
hyphenation or punctuation.

Usage: python make_docx.py <report.tex> <out.docx> <pandoc.exe> <earlier report: .txt or .docx> [figures]

With ``figures`` only the captions of new figures are highlighted, nothing else.
"""
import copy
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run

TEX, OUT, PANDOC, EARLIER = sys.argv[1:5]
ONLY_FIGURES = sys.argv[5:6] == ["figures"]
FONT = "Times New Roman"
BODY_STYLES = ("Body Text", "First Paragraph", "Compact", "Normal")
TEXT_WIDTH = 5.77          # inches between the margins


# ====================================================================== #
# what the earlier report said
# ====================================================================== #
REFERENCE = re.compile(r"(?i)\b(figures?|tables?|listings?|sections?)\s+\d+(\.\d+)*(\s+(to|and)\s+\d+(\.\d+)*)?")


def norm(text: str) -> str:
    """Letters and digits only; the numbers of figures, tables and sections are left out because they shift."""
    text = re.sub(r"\[lst:[\w-]+\]", "", text)          # a listing reference Pandoc has not numbered yet
    return re.sub(r"[^a-z0-9]", "", REFERENCE.sub(r"\1", text).lower())


def earlier_text(path: str) -> tuple[str, str, str]:
    """The earlier report: (running text without page numbers, everything), both normalised, and the raw text."""
    if path.lower().endswith(".docx"):
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return norm(" ".join(parts)), norm(" ".join(parts)), " ".join(parts)
    with open(path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    lines = [line for line in lines if not line.startswith("=== PAGE")]
    flow = [line for line in lines if not re.fullmatch(r"\s*\d{1,3}\s*", line)]     # page and listing line numbers
    return norm(" ".join(flow)), norm(" ".join(lines)), " ".join(lines)


BASE_FLOW, BASE_ALL, BASE_RAW = earlier_text(EARLIER)
# the table rows of the earlier report (only a .docx has them), so a changed cell can be found in its row
BASE_ROWS = ([[norm(cell.text) for cell in row.cells] for table in Document(EARLIER).tables for row in table.rows]
             if EARLIER.lower().endswith(".docx") else [])


def changed_cells(cells: list[str]) -> list[bool] | None:
    """Which cells of a row differ from the most similar row of the earlier report's tables (None: no such row)."""
    keys = [norm(text) for text in cells]
    best, score = None, 0
    for row in BASE_ROWS:
        if len(row) == len(keys) and row[0] == keys[0]:
            same = sum(1 for a, b in zip(row, keys) if a == b)
            if same > score:
                best, score = row, same
    if best is None or score * 2 < len(keys):
        return None
    return [a != b and bool(b) for a, b in zip(best, keys)]


def is_old(text: str) -> bool:
    key = norm(text)
    if not key:
        return True
    if len(key) < 12:           # a short piece must stand by itself there, not inside a longer word
        return re.search(r"(?<!\w)" + re.escape(text.strip()) + r"(?!\w)", BASE_RAW, re.IGNORECASE) is not None
    return key in BASE_FLOW or key in BASE_ALL


def sentences(text: str):
    """(start, end) of every sentence in ``text``."""
    spans, start = [], 0
    for m in re.finditer(r"(?<=[.!?;:])\s+(?=[A-Z(\[`'\"])", text):
        spans.append((start, m.start()))
        start = m.end()
    spans.append((start, len(text)))
    return [(a, b) for a, b in spans if text[a:b].strip()]


# ====================================================================== #
# run-level helpers
# ====================================================================== #
def run_elements(paragraph):
    return list(paragraph._p.iter(qn("w:r")))


def run_text(r) -> str:
    return "".join(t.text or "" for t in r.findall(qn("w:t")))


def paragraph_text(paragraph) -> str:
    return "".join(run_text(r) for r in run_elements(paragraph))


def highlight(paragraph, spans, figure: bool = False) -> None:
    """Yellow highlight on the character ranges ``spans`` of the paragraph's text."""
    if not spans or (ONLY_FIGURES and not figure):
        return
    pos = 0
    for r in run_elements(paragraph):
        texts = r.findall(qn("w:t"))
        length = len(run_text(r))
        start, end = pos, pos + length
        pos = end
        if not length:
            continue
        hits = sorted((max(a, start) - start, min(b, end) - start) for a, b in spans if a < end and b > start)
        if not hits:
            continue
        if len(texts) != 1 or hits == [(0, length)]:
            Run(r, paragraph).font.highlight_color = WD_COLOR_INDEX.YELLOW
            continue
        # split the run at every boundary: unmarked and marked pieces in turn
        text = texts[0].text
        merged = []
        for a, b in hits:
            if merged and not text[merged[-1][1]:a].strip():      # only white space between two marked pieces
                merged[-1] = (merged[-1][0], b)
            else:
                merged.append((a, b))
        pieces, at = [], 0
        for a, b in merged:
            pieces += [(text[at:a], False), (text[a:b], True)]
            at = b
        pieces.append((text[at:], False))
        anchor = r
        first = True
        for piece, marked in pieces:
            if not piece:
                continue
            target = r if first else copy.deepcopy(r)
            if not first:
                anchor.addnext(target)
            t = target.find(qn("w:t"))
            t.text = piece
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            rpr = target.find(qn("w:rPr"))
            if rpr is not None:
                for old in rpr.findall(qn("w:highlight")):
                    rpr.remove(old)
            if marked:
                Run(target, paragraph).font.highlight_color = WD_COLOR_INDEX.YELLOW
            anchor, first = target, False


def mark_new(paragraph, whole: bool = False, skip: int = 0) -> int:
    """Highlight what is new in a paragraph. ``whole`` judges it as one piece; ``skip`` leaves
    that many leading characters (an added number such as 'Table 4.1: ') out of the comparison."""
    text = paragraph_text(paragraph)
    body = text[skip:]
    if not body.strip():
        return 0
    spans = [(skip, len(text))] if whole else [(a + skip, b + skip) for a, b in sentences(body)]
    new = [(a, b) for a, b in spans if not is_old(text[a:b])]
    if whole and new and skip:
        new = [(0, len(text))]              # a new caption is new together with its number
    highlight(paragraph, new)
    return 0 if ONLY_FIGURES else len(new)


def set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def style_font(style, size, bold=None, italic=None, name=FONT):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(theme)) is not None:
            del fonts.attrib[qn(theme)]
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def field(paragraph, instruction: str, placeholder: str = "") -> None:
    """A complex field (TOC, PAGE) in the paragraph."""
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = paragraph.add_run()
        set_font(run, 12)
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = instruction
            run._r.append(el)
        elif kind == "text":
            run.text = placeholder
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def page_break_before(paragraph) -> None:
    paragraph.paragraph_format.page_break_before = True


def table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    tbl_pr.append(borders)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    for old in tbl_pr.findall(qn("w:jc")):
        tbl_pr.remove(old)
    tbl_pr.append(jc)


# ====================================================================== #
# convert
# ====================================================================== #
folder = os.path.dirname(os.path.abspath(TEX))
raw = OUT + ".raw.docx"
subprocess.run([PANDOC, os.path.basename(TEX), "-f", "latex", "-t", "docx", "-o", raw, "--resource-path=.",
                "--top-level-division=chapter", "--number-sections"], cwd=folder, check=True)
with open(TEX, encoding="utf-8") as fh:
    tex = fh.read()
listings = []       # (caption, label) in order
for options in re.findall(r"\\begin\{lstlisting\}\[(.*?)\]\n", tex):
    caption = re.search(r"caption=\{(.*)\},\s*label=", options)
    label = re.search(r"label=([\w:-]+)", options)
    text = caption.group(1) if caption else ""
    text = re.sub(r"\\texttt\{([^}]*)\}", r"\1", text).replace("\\_", "_")
    listings.append((text, label.group(1) if label else ""))

doc = Document(raw)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.left_margin, section.right_margin = Inches(1.5), Inches(1.0)
section.top_margin, section.bottom_margin = Inches(1.0), Inches(1.0)

for name in BODY_STYLES + ("Table", "Bibliography", "Block Text"):
    if name in [s.name for s in doc.styles]:
        style_font(doc.styles[name], 12)
for name, size, bold, italic in (("Heading 1", 16, True, False), ("Heading 2", 14, True, False),
                                 ("Heading 3", 12, True, False), ("Heading 4", 12, True, True),
                                 ("Image Caption", 10, False, False), ("Table Caption", 10, False, False),
                                 ("Caption", 10, False, False)):
    if name in [s.name for s in doc.styles]:
        style_font(doc.styles[name], size, bold, italic)

paragraphs = doc.paragraphs
first_heading = next(i for i, p in enumerate(paragraphs) if p.style.name == "Heading 1")
abstract_at = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "ABSTRACT")
toc_at = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "TABLE OF CONTENTS")
references_at = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "REFERENCES")
COVER_SIZES = {"BCSE307P COMPILER DESIGN LAB PROJECT": (16, 30), "(Slot": (14, 6), "AUTOMATED LLM": (16, 36),
               "Submitted by": (12, 42), "Under the guidance": (12, 30), "Prof.": (12, 6),
               "School of Computer": (14, 36), "October": (12, 6)}

chapter = figure_no = table_no = listing_no = 0
listing_index = 0
listing_numbers = {}
marked = 0
for i, p in enumerate(paragraphs):
    style = p.style.name
    text = p.text.strip()
    fmt = p.paragraph_format
    # ------------------------------------------------------------ cover page
    if i < abstract_at:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.line_spacing = 1.15
        for key, (size, before) in COVER_SIZES.items():
            if text.startswith(key):
                fmt.space_before = Pt(before)
                for run in p.runs:
                    set_font(run, size, bold=key not in ("Prof.", "October"))
        continue
    # ------------------------------------------------------------ front-matter titles
    if i in (abstract_at, toc_at):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_break_before(p)
        fmt.space_after = Pt(12)
        for run in p.runs:
            set_font(run, 16, bold=True)
        if i == toc_at:
            toc = copy.deepcopy(paragraphs[abstract_at + 1]._p)
            p._p.addnext(toc)
            toc_paragraph = type(p)(toc, p._parent)
            for r in list(toc_paragraph._p):
                if r.tag != qn("w:pPr"):
                    toc_paragraph._p.remove(r)
            toc_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            field(toc_paragraph, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
        continue
    if i == references_at:
        p.style = doc.styles["Heading 1"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_break_before(p)
        fmt.space_after = Pt(12)
        for run in p.runs:
            set_font(run, 16, bold=True, color=RGBColor(0, 0, 0))
        continue
    # ------------------------------------------------------------ headings
    if style == "Heading 1":
        chapter += 1
        figure_no = table_no = listing_no = 0
        runs = p.runs
        number = runs[0].text.strip()
        runs[0].text = f"CHAPTER {number}"
        runs[0].add_break(WD_BREAK.LINE)
        for run in runs[1:]:
            run.text = run.text.replace("\t", "")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_break_before(p)
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(18), 1.5
        for run in p.runs:
            set_font(run, 16, bold=True, color=RGBColor(0, 0, 0))
        marked += mark_new(p, whole=True, skip=len(f"CHAPTER {number}"))
        continue
    if style in ("Heading 2", "Heading 3", "Heading 4"):
        size = 14 if style == "Heading 2" else 12
        for run in p.runs:
            run.text = run.text.replace("\t", "  ")
            mono = run.style is not None and run.style.name == "Verbatim Char"
            set_font(run, size, bold=True, italic=style == "Heading 4", color=RGBColor(0, 0, 0),
                     name="Consolas" if mono else FONT)
            if style == "Heading 2":
                run.font.all_caps = True
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(12), Pt(6), 1.5
        number = p.runs[0].text if p.runs else ""
        marked += mark_new(p, whole=True, skip=len(number) + 2)
        continue
    # ------------------------------------------------------------ captions
    if style in ("Image Caption", "Table Caption"):
        if style == "Image Caption":
            figure_no += 1
            prefix = f"Figure {chapter}.{figure_no}: "
        else:
            table_no += 1
            prefix = f"Table {chapter}.{table_no}: "
        new = not is_old(paragraph_text(p))
        first = p.runs[0] if p.runs else p.add_run()
        first.text = prefix + first.text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.line_spacing = 1.0
        fmt.space_before, fmt.space_after = Pt(4), Pt(8)
        for run in p.runs:
            set_font(run, 10, italic=False)
        if new and (style == "Image Caption" or not ONLY_FIGURES):
            highlight(p, [(0, len(paragraph_text(p)))], figure=style == "Image Caption")
            marked += 1
        continue
    if style == "Captioned Figure":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.keep_with_next = True
        continue
    # ------------------------------------------------------------ code
    if style == "Source Code":
        listing_no += 1
        caption, label = listings[listing_index] if listing_index < len(listings) else ("", "")
        listing_index += 1
        listing_numbers[label] = f"{chapter}.{listing_no}"
        fmt.line_spacing = 1.0
        fmt.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ppr = p._p.get_or_add_pPr()
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:color"), "auto")
        shade.set(qn("w:fill"), "F3F3F3")
        ppr.append(shade)
        for run in p.runs:
            set_font(run, 8.5, name="Consolas")
        code_lines = [line for line in p.text.split("\n") if line.strip()]
        fresh = sum(1 for line in code_lines if not is_old(line))
        code_new = fresh * 2 > len(code_lines)
        if code_new:
            highlight(p, [(0, len(paragraph_text(p)))])
            marked += 1
        if caption:
            cap = copy.deepcopy(paragraphs[abstract_at + 1]._p)
            p._p.addprevious(cap)
            cap_paragraph = type(p)(cap, p._parent)
            for child in list(cap_paragraph._p):
                if child.tag != qn("w:pPr"):
                    cap_paragraph._p.remove(child)
            run = cap_paragraph.add_run(f"Listing {chapter}.{listing_no}: {caption}")
            set_font(run, 10)
            cap_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_paragraph.paragraph_format.line_spacing = 1.0
            cap_paragraph.paragraph_format.keep_with_next = True
            if not is_old(caption) and not ONLY_FIGURES:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        continue
    # ------------------------------------------------------------ body text
    if style in BODY_STYLES:
        for run in p.runs:
            mono = run.style is not None and run.style.name == "Verbatim Char"
            set_font(run, 11 if mono else 12, name="Consolas" if mono else FONT)
        if i > references_at:
            fmt.line_spacing = 1.15
            if text.endswith(":") and len(text) < 40:           # a group heading: Journals, Conferences, ...
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(10), Pt(4), True
                for run in p.runs:
                    set_font(run, 12, bold=True)
                marked += mark_new(p, whole=True)
                continue
            fmt.left_indent, fmt.first_line_indent = Inches(0.5), Inches(-0.5)
            # a long address cannot be broken, and justifying the line before it stretches the words apart
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if "http" in text else WD_ALIGN_PARAGRAPH.JUSTIFY
            marked += mark_new(p, whole=True)
            continue
        fmt.line_spacing = 1.15 if abstract_at < i < toc_at else 1.5
        fmt.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        marked += mark_new(p)

# listing references that Pandoc left as [lst:...]
for p in doc.paragraphs:
    if "[lst:" in p.text:
        for r in run_elements(p):
            for t in r.findall(qn("w:t")):
                if t.text and "[lst:" in t.text:
                    t.text = re.sub(r"\[(lst:[\w-]+)\]", lambda m: listing_numbers.get(m.group(1), m.group(0)), t.text)

# ------------------------------------------------------------------ tables
for index, table in enumerate(doc.tables):
    table_borders(table)
    cover = index == 0
    # a table whose caption is new is new as a whole; a table the earlier report had is judged cell by cell
    before = table._tbl.getprevious()
    while before is not None and before.tag != qn("w:p"):
        before = before.getprevious()
    caption = ""
    if before is not None and any(s.get(qn("w:val")) == "TableCaption" for s in before.iter(qn("w:pStyle"))):
        caption = "".join(t.text or "" for t in before.iter(qn("w:t")))
    all_new = not cover and bool(caption) and not is_old(caption.split(": ", 1)[-1])
    for r, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))          # a row never breaks across two pages
        changes = None if all_new or cover else changed_cells([cell.text for cell in row.cells])
        if changes is not None:
            changes = [False] * len(row.cells)                # a row the earlier report had: a corrected cell stays plain
        if changes is None and BASE_ROWS and not (all_new or cover):
            if not is_old(" ".join(cell.text for cell in row.cells)):
                changes = [True] * len(row.cells)             # a row the earlier table did not have
        for c, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    mono = run.style is not None and run.style.name == "Verbatim Char"
                    set_font(run, 10 if mono else 11, name="Consolas" if mono else FONT)
                if (all_new or (changes is not None and changes[c])) and paragraph_text(p).strip():
                    highlight(p, [(0, len(paragraph_text(p)))])
                    marked += 1
                elif changes is not None:
                    continue                                  # the same as in the earlier table
                elif not cover:
                    marked += mark_new(p, whole=True)

# ------------------------------------------------------------------ pictures
for shape in doc.inline_shapes:
    limit = Inches(2.9) if shape is doc.inline_shapes[0] else Inches(TEXT_WIDTH)
    if shape.width > limit or shape is doc.inline_shapes[0]:
        ratio = limit / shape.width
        shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)
    if shape.height > Inches(7.6):
        ratio = Inches(7.6) / shape.height
        shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)

# ------------------------------------------------------------------ page numbers, centred at the bottom
footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(footer, "PAGE", "1")
section.different_first_page_header_footer = True

doc.save(OUT)
os.remove(raw)
print(f"saved {OUT}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.inline_shapes)} pictures; "
      f"{marked} pieces marked as new")
